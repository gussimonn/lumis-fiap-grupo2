import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

os.chdir(r"C:\Users\gusta\Downloads\LumisOS\00_Lumis\Design_System")

NAVY = RGBColor(0x0F, 0x2D, 0x3A); INK = RGBColor(0x18, 0x25, 0x2C); MUTED = RGBColor(0x6B, 0x7C, 0x85)  # Profundo, Tinta, Névoa
AMBER_TXT = RGBColor(0xB7, 0x79, 0x1F); TEAL = RGBColor(0x1E, 0x8C, 0x6B); RED = RGBColor(0xB4, 0x44, 0x2F)  # TEAL = Vital escuro
SERIF = "Georgia"; SANS = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.2); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
sec.different_first_page_header_footer = True


def font(style, name, size, color=None, bold=False, italic=False):
    f = style.font; f.name = name; f.size = Pt(size); f.bold = bold; f.italic = italic
    if color is not None:
        f.color.rgb = color
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def para(style, before=0, after=6, line=1.15, keep=False):
    pf = style.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if keep:
        pf.keep_with_next = True


def border(el, side, color, sz=12, space=8):
    pPr = el.get_or_add_pPr(); bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr'); pPr.append(bdr)
    e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); bdr.append(e)


def shade(pr, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); pr.append(s)


st = doc.styles
font(st['Normal'], SANS, 11, INK); para(st['Normal'], 0, 6, 1.15)
font(st['Title'], SERIF, 28, NAVY); para(st['Title'], 0, 6, 1.0)
tp = st['Title'].element.get_or_add_pPr()
for b in tp.findall(qn('w:pBdr')):
    tp.remove(b)
font(st['Subtitle'], SANS, 13, MUTED); para(st['Subtitle'], 0, 18, 1.15)
for e in st['Subtitle'].element.get_or_add_rPr().findall(qn('w:spacing')):
    e.getparent().remove(e)
font(st['Heading 1'], SERIF, 16, NAVY); para(st['Heading 1'], 18, 6, 1.1, True)
border(st['Heading 1'].element, 'bottom', '3DBE93', 8, 4)
font(st['Heading 2'], SERIF, 12.5, NAVY); para(st['Heading 2'], 12, 4, 1.1, True)
font(st['Heading 3'], SANS, 11, NAVY, bold=True); para(st['Heading 3'], 8, 2, 1.1, True)
font(st['Caption'], SANS, 9, MUTED, italic=True); para(st['Caption'], 4, 10, 1.0)


def new_pstyle(name, base='Normal'):
    s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH); s.base_style = st[base]; s.quick_style = True; return s


def new_cstyle(name):
    s = st.add_style(name, WD_STYLE_TYPE.CHARACTER); s.quick_style = True; return s


d = new_pstyle('Lumis Destaque'); font(d, SERIF, 12, NAVY, italic=True); para(d, 6, 10, 1.25)
d.paragraph_format.left_indent = Cm(0.4); border(d.element, 'left', '3DBE93', 24, 10)
shade(d.element.get_or_add_pPr(), 'EEF7F3')
r = new_pstyle('Lumis Rótulo'); font(r, SANS, 9, TEAL, bold=True); para(r, 0, 2, 1.0)
rpr = r.element.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '30'); rpr.append(sp)
rpr.append(OxmlElement('w:caps'))
t = new_pstyle('Lumis Tabela'); font(t, SANS, 9.5, INK); para(t, 2, 2, 1.0)
n = new_pstyle('Lumis Nota'); font(n, SANS, 9, MUTED); para(n, 0, 4, 1.1)
for name, col in (('Tag Fonte', TEAL), ('Tag Hipótese', AMBER_TXT), ('Tag Não consta', MUTED), ('Tag Risco', RED)):
    c = new_cstyle(name); font(c, SANS, 9, col, bold=True)

# Capa
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture('artefato/feixe_capa.jpg', width=Cm(16))
doc.add_paragraph('Fase X · Entrega Y', style='Lumis Rótulo').paragraph_format.space_before = Pt(28)
doc.add_paragraph('Título da entrega', style='Title')
doc.add_paragraph('Subtítulo: a pergunta que o documento responde, em uma linha', style='Subtitle')
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(110)
p.add_run().add_picture('logo/lumis-insight_horizontal.png', width=Cm(5.2))
for line in ('Gestão em IA · FIAP', 'Equipe: Nome Sobrenome, Nome Sobrenome, Nome Sobrenome', 'Data: dd/mm/aaaa'):
    doc.add_paragraph(line, style='Lumis Nota')
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# Cabeçalho e rodapé das páginas internas
h = sec.header.paragraphs[0]; h.text = ''
h.add_run().add_picture('logo/lumis-insight_horizontal.png', width=Cm(3.2))
hr = h.add_run('\tFX-EY · Nome da entrega'); hr.font.size = Pt(8.5); hr.font.color.rgb = MUTED; hr.font.name = SANS
hp = st['Header'].element.get_or_add_pPr()
for tabs in hp.findall(qn('w:tabs')):
    hp.remove(tabs)
h.paragraph_format.tab_stops.add_tab_stop(Cm(16), 2)  # 2 = direita
border(h._p, 'bottom', '0F2D3A', 4, 4)
f = sec.footer.paragraphs[0]; f.text = ''
rr = f.add_run('Lumis Intelligence · Gestão em IA (FIAP) · página '); rr.font.size = Pt(8.5); rr.font.color.rgb = MUTED
run = f.add_run(); run.font.size = Pt(8.5); run.font.color.rgb = MUTED
for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if kind:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
    run._r.append(e)
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Corpo de exemplo
doc.add_paragraph('Em uma frase', style='Lumis Rótulo')
doc.add_paragraph('A conclusão abre o documento. Escreva aqui, em duas ou três linhas, o que o leitor precisa saber se parar de ler neste ponto.', style='Lumis Destaque')
doc.add_heading('1. Título de seção', level=1)
p = doc.add_paragraph('Texto corrido em Calibri 11, cor tinta, entrelinha 1,15. O dado que sustenta uma decisão leva a origem logo depois da frase ')
p.add_run('[fonte: Cap. 2, Quadro 9]').style = st['Tag Fonte']
p.add_run('. Leitura da equipe sem quadro que a sustente fica marcada como ')
p.add_run('[hipótese]').style = st['Tag Hipótese']
p.add_run(', e o que não está no material aparece como ')
p.add_run('[não consta]').style = st['Tag Não consta']
p.add_run('. Um ponto que contradiz um compromisso vigente leva ')
p.add_run('[risco: C2]').style = st['Tag Risco']
p.add_run('.')
doc.add_heading('1.1 Subseção', level=2)
doc.add_paragraph('Subseções em Georgia 12,5. No corpo, usar no máximo dois níveis. O terceiro, em Calibri negrito, serve para rotular tabelas e listas longas.')
doc.add_heading('Tabela padrão', level=3)
rows = [('Coluna A', 'Coluna B', 'Coluna C'), ('Item', 'Valor', '[fonte: ...]'), ('Item', 'Valor', '[hipótese]'), ('Item', 'Valor', '[não consta]')]
tb = doc.add_table(rows=len(rows), cols=3); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
bd = OxmlElement('w:tblBorders')
for side in ('top', 'bottom', 'insideH'):
    e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB'); bd.append(e)
tb._tbl.tblPr.append(bd)
for i, row in enumerate(rows):
    for j, txt in enumerate(row):
        c = tb.cell(i, j); c.text = ''; pp = c.paragraphs[0]; pp.style = st['Lumis Tabela']; rn = pp.add_run(txt)
        tcPr = c._tc.get_or_add_tcPr()
        if i == 0:
            rn.bold = True; rn.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(tcPr, '0F2D3A')
        elif i % 2 == 0:
            shade(tcPr, 'F1F6F4')
doc.add_paragraph('Tabela 1. Legenda curta abaixo da tabela, em Calibri 9 itálico, com a fonte dos dados no fim.', style='Caption')
doc.add_heading('2. Fechamento', level=1)
doc.add_paragraph('Próximos passos e decisões registradas. A conclusão já abriu o documento e não se repete aqui.')
import sys
doc.save(sys.argv[1] if len(sys.argv) > 1 else 'Modelo_Entrega_Lumis.docx')
print('ok')
