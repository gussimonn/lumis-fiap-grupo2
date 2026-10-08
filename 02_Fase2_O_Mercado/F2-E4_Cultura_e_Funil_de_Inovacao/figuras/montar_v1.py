"""F2-E4 v1: Diagnóstico de Cultura e Funil de Inovação, versão enxuta para o documento integrado.
Base: F2-E4_Levantamento_v1.md e decisões D-037 a D-040. Cargos citados pelo papel, sem nome.
Formatação: a mesma da F2-E3 v1 (capa, cabeçalho e rodapé do design system, a partir da F2-E1 v2)."""
import re, sys, docx
from docx.shared import Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = 'C:/Users/gusta/Downloads/LumisOS/'
MODELO = BASE + '02_Fase2_O_Mercado/F2-E1_Mapa_do_Territorio/F2-E1_Mapa_do_Territorio_v2.docx'
FIG = BASE + '02_Fase2_O_Mercado/F2-E4_Cultura_e_Funil_de_Inovacao/figuras/'
OUT = sys.argv[1]

d = docx.Document(MODELO)
st = d.styles
body = d.element.body

hp = d.sections[0].header.paragraphs[0]
feito = False
for r in hp.runs:
    if r.text.strip():
        r.text = '' if feito else '\tF2-E4 · Cultura e Funil de Inovação'
        feito = True

for p, t in ((d.paragraphs[0], 'Cultura e Funil de Inovação'),
             (d.paragraphs[1], 'O que separa o técnico do comercial, como decidir o que se constrói e o que recusamos')):
    p.runs[0].text = t
    for r in p.runs[1:]:
        r.text = ''
sub = d.paragraphs[1]._p
apagar = False
for el in list(body):
    if apagar and el.tag != qn('w:sectPr'):
        body.remove(el)
    if el is sub:
        apagar = True

TAG = re.compile(r'(\[(?:fonte|hipótese|não consta|risco)[^\]]*\])')
TAGSTYLE = {'fonte': 'Tag Fonte', 'hipótese': 'Tag Hipótese', 'não consta': 'Tag Não consta', 'risco': 'Tag Risco'}


def runs(p, text):
    for i, seg in enumerate(text.split('**')):
        for part in TAG.split(seg):
            if not part:
                continue
            r = p.add_run(part)
            if TAG.fullmatch(part):
                r.style = st[TAGSTYLE[next(k for k in TAGSTYLE if part[1:].startswith(k))]]
            elif i % 2:
                r.bold = True
    return p


def P(text, style='Normal'):
    return runs(d.add_paragraph(style=style), text)


def B(text):
    return P(text, 'List Bullet')


def H1(t): d.add_heading(t, level=1)


def FIGURA(arq, legenda):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(FIG + arq, width=Cm(16))
    P(legenda, 'Caption')


def shade(cell, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(s)


def TABELA(linhas, larguras, legenda):
    tb = d.add_table(rows=len(linhas), cols=len(linhas[0])); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB')
        bd.append(e)
    tb._tbl.tblPr.append(bd); tb.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(th)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in tb.rows:
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i, row in enumerate(linhas):
        for j, txt in enumerate(row):
            c = tb.cell(i, j); c.width = Cm(larguras[j])
            p = c.paragraphs[0]; p.style = st['Lumis Tabela']
            p.paragraph_format.keep_with_next = True
            if i == 0:
                r = p.add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, '0F2D3A')
            else:
                runs(p, txt)
                if i % 2 == 0:
                    shade(c, 'F1F6F4')
    P(legenda, 'Caption')


# ---------------- abertura ----------------
P('O técnico e o comercial discordam sobre quanto tempo validar, mas partem da mesma crença: a qualidade se prova antes de o produto ir a campo, pela média. Propomos um único critério para decidir o que se promete e o que se constrói, o desempenho medido em campo por grupo de pacientes, e com ele recusamos três dos sete pedidos do backlog.', 'Lumis Destaque')
P('O enunciado trata a fratura interna e o excesso de oportunidades como um problema só: falta um critério que as áreas aceitem. Os dados mostram onde ele falta.')
B('**Na promessa.** Só 22% do time técnico concorda que a Lumis promete ao cliente o que o produto entrega. No comercial, são 79% [fonte: Cap. 2, Quadro 15].')
B('**No backlog.** Os pedidos somam 97 meses-pessoa, mais de três vezes os 30 disponíveis no semestre, e estão organizados por quem pediu, sem critério de risco, de prova ou de encerramento [fonte: Cap. 2, Quadro 16].')
B('**Na escalada.** Em todos os grupos, só 21% a 29% sabem a quem levar um problema ético do produto [fonte: Cap. 2, Quadro 15].')

# ---------------- 1. cultura ----------------
H1('1. A cultura em três camadas')
P('Usamos as três camadas de Schein¹: o que se vê, o que se diz e o que o grupo dá como óbvio sem perceber. O técnico vê a promessa longe da entrega e se diz pouco ouvido; o comercial vê o contrário. A rotatividade em doze meses foi de 19% na empresa e de 27% no time de dados [fonte: Cap. 2, Quadro 15].')
FIGURA('v1_fig1_clima.png', 'Figura 1. Concordância com cada afirmação, por grupo. Pesquisa de julho de 2026, 81 respostas de 96 colaboradores. Fonte: Cap. 2, Quadro 15.')
TABELA([
    ['Camada', 'O que encontramos'],
    ['Artefatos', 'O número de vitrine (94% de acurácia) vem de uma validação de 2023 e continua no material; em campo, o resultado é 87,6%. Um erro corrigido no sistema em 01/2026 seguiu no material comercial. O erro por grupo foi medido pelo hospital, e as reclamações nunca foram cruzadas com o desempenho. O backlog registra quem pediu, e não o risco.'],
    ['Valores declarados', 'O técnico diz que o comercial promete o que o produto não entrega. O comercial diz que o técnico atrasa tudo em nome de um rigor que o cliente não pede.'],
    ['Pressupostos', 'Cada área acredita vender uma coisa: confiabilidade ou velocidade. Por baixo das duas, a mesma crença: "aqui, modelo validado é modelo bom".'],
], [3.2, 12.4], 'Tabela 1. A cultura da Lumis nas três camadas. Fonte: Cap. 2, seções 1.2 e 2.5 e Quadros 9, 10, 11, 14, 16 e 17; leitura da equipe.')
P('**O pressuposto profundo** [hipótese]: "Aqui, modelo validado é modelo bom: a qualidade se prova antes de o produto ir a campo, pela média, e depois disso o número vale." O capítulo descreve a crença de cada área, confiabilidade para uma e velocidade para a outra [fonte: Cap. 2, 2.5], e as duas aceitam essa premissa. Se a qualidade se resolve antes, só resta discutir quanto tempo gastar antes, e sem número de campo por grupo a discussão não tem árbitro. Por isso ela nunca chegou "para uma sala onde ela poderia ser resolvida" [fonte: Cap. 2, 1.2].')
P('A crença comum explica o que a leitura de cada área não explica: por que as duas concordam em não saber a quem escalar, por que nenhuma mediu o erro por grupo antes do hospital e por que o comercial tem mais voz (63% contra 34%). Depois do lançamento, a única informação nova que circula é a do cliente que compra, e quem a traz é o comercial [hipótese]. Inferimos o pressuposto de quadros, sem entrevistas; a mesa de campo e a pesquisa de 01/2027 servem para testá-lo.')

# ---------------- 2. intervenções ----------------
H1('2. Duas intervenções')
P('Comunicado interno não resolve pressuposto [fonte: Cap. 2, 2.5]. As duas intervenções mudam o que se mede e quem decide.')
TABELA([
    ['Intervenção', 'Responsável', 'Prazo', 'Sinal de que funcionou'],
    ['**A. Regra única de prova.** Nenhum número de desempenho sai da empresa sem o resultado de campo aberto por grupo, com amostra e data. Vale para material, proposta, aditivo e liberação de versão. Incidente só fecha com o material corrigido.',
     'Head of AI Management; o Responsável Comercial aplica',
     '30 dias para as cinco afirmações do painel',
     'Afirmações públicas auditáveis: de 0 para 5 de 5, ou retiradas'],
    ['**B. Mesa de campo quinzenal.** Técnico, comercial, dados, produto e DPO leem juntos o erro por grupo e as reclamações por faixa de CEP. A mesa propõe o prazo de validação, que o CTO decide, e responde com data aos pedidos do comercial.',
     'Responsável por Produto',
     'Primeira reunião em 15 dias; prazos de validação publicados em 60',
     'Reclamações cruzadas com o desempenho: de nunca para toda quinzena'],
], [7.0, 3.0, 2.8, 2.8], 'Tabela 2. As duas intervenções. Responsáveis, prazos e metas são proposta da equipe. Linhas de base: Cap. 2, Quadros 11 e 17; Entrega 2.')
P('A regra produz o número por grupo, e a mesa é onde ele vira decisão de produto e de venda. A mesa usa o horário da leitura quinzenal que o compromisso de Não Amplificação de Danos já prevê, e as duas áreas ganham algo: o técnico, prazo de validação com regra; o comercial, prazo conhecido e um número que resiste a uma due diligence. Para medir a percepção, repetimos as afirmações da Figura 1 numa pesquisa curta em 01/2027, aplicada por terceiro, e esperamos pelo menos 10 pontos a mais no técnico e em "sei a quem escalar" (hoje cerca de 26% de quem respondeu) [fonte: Cap. 2, Quadro 15]. O resultado é julgado na pesquisa de 07/2027.')

# ---------------- 3. funil ----------------
H1('3. Funil de inovação')
P('Cada passagem é um portão com quatro saídas: seguir, encerrar, pausar com data ou voltar, como no Stage-Gate². Os critérios são eliminatórios, sem nota somada, para que receita não compense risco aos pacientes. O portão segue a Entrega 3: o Responsável por Produto leva o pedido, a CEO decide e o Head of AI Management e a DPO dão parecer.')
FIGURA('v1_fig2_funil.png', 'Figura 2. As cinco etapas, o que é preciso para avançar e o que o comercial pode dizer em cada uma. Proposta da equipe.')
TABELA([
    ['Etapa', 'Entra quando', 'Avança quando', 'Encerra quando'],
    ['Entrada', 'Ficha com problema, cliente nomeado, base legal dos dados e métrica de sucesso', 'Cliente no foco da Entrega 1 ou exceção registrada; nenhuma dependência nova de fornecedor único; nada contra os compromissos; responsável com cargo', 'Falha em qualquer critério'],
    ['Descoberta (até 10% do esforço)', 'Aprovada na Entrada', 'Cliente com intenção de compra e dados autorizados por escrito; amostra para medir o erro em todo grupo; responsável do domínio nomeado', 'A fatia ou o prazo acabam e falta algum item'],
    ['Modo sombra (até mais 30%)', 'Roda em paralelo, sem decidir', 'Quatro leituras quinzenais com o pior grupo até 1,5 vez o erro do melhor (no crédito, aprovação por faixa de CEP de pelo menos 0,8 da melhor)', 'Seis leituras com algum grupo acima de 2 vezes'],
    ['Produção limitada', 'Condições da Entrega 5; autorização da CEO por cliente', 'Leitura sem piora em nenhum grupo; número publicado', 'Sem correção em seis leituras'],
    ['Escala', 'Auditoria independente contratada', 'Indicadores contínuos', 'Suspensão pela Entrega 3'],
], [3.0, 3.7, 5.6, 3.3], 'Tabela 3. Critérios de entrada, avanço e encerramento. Fatias e número de leituras são proposta da equipe; os limites de erro repetem as Entregas 2, 3 e 5.')
B('**Quem pediu não decide.** Os pedidos do sócio-fundador, do investidor e do comercial passam pelo mesmo portão. A CEO só decide contra um parecer contrário por escrito, e o caso vai ao conselho.')
B('**Recurso por etapa e encerramento definido antes.** Nenhuma iniciativa passa de 10 meses-pessoa sem nova decisão. Encerra a que gastar a fatia sem evidência nova, não obtiver o direito de uso dos dados até a data marcada ou tiver a estimativa de esforço aumentada em mais de 50%. Pipeline, explicabilidade e ISO/IEC 42001 entram pelo mesmo portão e recebem recurso por marco de entrega.')

# ---------------- 4. distribuição ----------------
H1('4. Distribuição dos 30 meses-pessoa')
P('Proporção fixa não funciona aqui: na divisão 70/20/10 de Nagji e Tuff³, a faixa adjacente teria 6 meses-pessoa, e a menor iniciativa adjacente pede 11. Financiamos etapas, e a proporção resulta da ordem de prioridade: compromisso já assumido, depois o ativo que a Entrega 1 aponta, depois a receita por esforço.')
FIGURA('v1_fig3_capacidade.png', 'Figura 3. O backlog contra a capacidade do semestre. Fonte: Cap. 2, Quadro 16; distribuição proposta pela equipe.')
TABELA([
    ['Horizonte', 'O que recebe', 'Meses-pessoa'],
    ['Melhoria do produto atual', 'Compromissos já assumidos 3,0; correção do viés 3,5; explicabilidade 9,0; primeira fatia do pipeline 7,2; diagnóstico da ISO/IEC 42001 0,7', '23,4 (78%)'],
    ['Expansão adjacente', 'Crédito para bancos, até o modo sombra', '5,6 (19%)'],
    ['Transformação', 'Descoberta da prova de desempenho por grupo como oferta', '1,0 (3%)'],
], [3.6, 9.6, 2.4], 'Tabela 4. Distribuição proposta. Proposta da equipe sobre o Cap. 2, Quadro 16.')
B('**Melhoria.** O erro de 31,8% e os compromissos das Entregas 2, 3 e 5 não têm linha no backlog, e expandir antes de resolver a base de dados é "replicação de um problema em escala maior" [fonte: Cap. 2, 2.6]. Não sabemos se a capacidade já desconta esse trabalho [não consta] e tratamos como se não descontasse.')
B('**Adjacência.** O crédito tem a maior receita por esforço, R$ 0,6 mi por mês-pessoa [fonte: conta da equipe sobre o Quadro 16], 5 bancos já clientes e um terreno onde a Aster não chega. Fica fora do foco da Entrega 1 e entra como exceção registrada. Pela Entrega 5, só vai à produção depois das condições.')
B('**Transformação.** Um mês-pessoa para testar se a prova de desempenho por grupo, o ativo que a Entrega 1 manda construir, vira oferta. É a nossa resposta à plataforma que o Vetor propõe [fonte: Cap. 2, 1.1].')

# ---------------- 5. recusa ----------------
H1('5. O que recusamos')
TABELA([
    ['Iniciativa', 'Quem pediu', 'Decisão', 'Por quê'],
    ['Módulo veterinário', 'Sócio-fundador', '**Recusado**', 'Domínio cujos dados e erros a Lumis não conhece, sem responsável do domínio, fora do foco. Ocuparia 53% da capacidade por 11% da receita listada'],
    ['Agente de triagem', 'Comercial', '**Recusado neste ciclo**', 'Automatizaria de novo a etapa do erro de 31,8%. Volta ao portão com o pior grupo até 1,5 vez o melhor e a explicabilidade pronta'],
    ['México', 'Investidor', '**Fora desta janela**', 'Ocuparia 73% da capacidade, sem dado local para medir o erro por grupo. Reavaliado ao fim do semestre'],
    ['Crédito para bancos', 'Comercial', 'Até o modo sombra', 'Seção 4'],
    ['Reescrita do pipeline', 'Time técnico', 'Primeira fatia', 'Rastreio da origem dos dados e erro por grupo automático'],
    ['Explicabilidade', 'Regulador e clientes', 'Por marco', 'Dá conteúdo à revisão humana'],
    ['ISO/IEC 42001', 'Jurídico', 'Só o diagnóstico', 'Certificação no ciclo seguinte'],
], [3.2, 2.6, 2.8, 7.0], 'Tabela 5. Decisão sobre cada pedido do backlog. Fonte: Cap. 2, Quadro 16; decisões da equipe.')
P('As três recusas, uma de cada lado (sócio-fundador, comercial e investidor), deixam de lado R$ 11,3 mi dos R$ 19,7 mi de receita potencial listada [fonte: conta da equipe sobre o Quadro 16], um teto sem prazo garantido. O veterinário só volta com parceiro de dados autorizados, responsável do domínio e o erro por grupo dentro do limite no Brasil.')

# ---------------- 6. coerência ----------------
H1('6. Coerência com o que já assumimos')
P('A regra de prova põe em prática a Transparência da Declaração de Intenção da Fase 1. A correção do viés e o modo sombra cumprem a Não Amplificação de Danos, a revisão de casos contestados em até 48 horas não muda, e cada decisão do funil fica registrada para a Responsabilidade e Prestação de Contas. Os cargos são os da Entrega 3. Vão para o memorando dois riscos: triplicar o time técnico, como propõe o Vetor, antes de as intervenções rodarem leva gente nova para o mesmo pressuposto, e por isso propomos contratar em ondas; e o Responsável por Dados aparece em quase todos os portões, num time com rotatividade de 27% [hipótese].')

H1('Referências')
P('1. SCHEIN, E. H.; SCHEIN, P. A. Organizational Culture and Leadership. 5. ed. Hoboken: Wiley, 2017. A bibliografia do curso registra a edição como de 2016.', 'Lumis Nota')
P('2. COOPER, R. G. Stage-gate systems: a new tool for managing new products. Business Horizons, v. 33, n. 3, p. 44-54, 1990.', 'Lumis Nota')
P('3. NAGJI, B.; TUFF, G. Managing Your Innovation Portfolio. Harvard Business Review, maio 2012. Disponível em: https://hbr.org/2012/05/managing-your-innovation-portfolio.', 'Lumis Nota')
P('Dados da Lumis: Cap. 2, Anexo A. Classificações, critérios, metas, prazos e cargos propostos são da equipe. A IA generativa (Claude Code, Anthropic) apoiou a pesquisa e a redação, com revisão da equipe.', 'Lumis Nota')

d.save(OUT)
print('ok')
