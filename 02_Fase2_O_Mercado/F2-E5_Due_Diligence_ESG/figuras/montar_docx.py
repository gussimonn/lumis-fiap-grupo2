"""F2-E5 v1: Due Diligence ESG da Expansão.
Conteúdo: F2-E5_Materialidade_v1.md, F2-E5_Equidade_v1.md, F2-E5_Privacidade_v1.md e F2-E5_Metrica_Publica_v1.md.
Base de formatação: F2-E1 v2 (mesma base da F2-E2 v1), estilos do Modelo_Entrega_Lumis.
Uso: python montar_docx.py [saida.docx]"""
import os, re, sys, docx
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

AQUI = os.path.dirname(os.path.abspath(__file__))
FASE = os.path.dirname(os.path.dirname(AQUI))
MODELO = os.path.join(FASE, 'F2-E1_Mapa_do_Territorio', 'F2-E1_Mapa_do_Territorio_v2.docx')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(AQUI), 'F2-E5_Due_Diligence_ESG_v1.docx')

d = docx.Document(MODELO)
st = d.styles
body = d.element.body

# cabeçalho das páginas internas
runs_h = d.sections[0].header.paragraphs[0].runs
runs_h[1].text = '\tF2-E5 · Due Diligence ESG'
runs_h[2].text = ''

# título e subtítulo; apaga o resto do corpo
d.paragraphs[0].runs[0].text = 'Due Diligence ESG da Expansão'
for r in d.paragraphs[0].runs[1:]:
    r.text = ''
d.paragraphs[1].runs[0].text = ('Relatório ao Comitê de Investimento: o que a expansão leva junto '
                                'e o que precisa estar resolvido antes')
for r in d.paragraphs[1].runs[1:]:
    r.text = ''
sub = d.paragraphs[1]._p
apagar = False
for el in list(body):
    if apagar and el.tag != qn('w:sectPr'):
        body.remove(el)
    if el is sub:
        apagar = True

TAG = re.compile(r'(\[(?:fonte|hipótese|não consta|risco)[^\]]*\])')
TAGSTYLE = {'fonte': 'Tag Fonte', 'hipótese': 'Tag Hipótese', 'não consta': 'Tag Não consta', 'risco': 'Tag Risco'}


def runs(p, text):
    for i, seg in enumerate(text.split('**')):
        for part in TAG.split(seg):
            if not part:
                continue
            r = p.add_run(part)
            if TAG.fullmatch(part):
                r.style = st[TAGSTYLE[next(k for k in TAGSTYLE if part[1:].startswith(k))]]
            elif i % 2:
                r.bold = True
    return p


def P(text, style='Normal'):
    return runs(d.add_paragraph(style=style), text)


def B(text): return P(text, 'List Bullet')
def N(text): return P(text, 'List Number')
def H1(t): d.add_heading(t, level=1)
def DESTAQUE(texto): return P(texto, 'Lumis Destaque')


def FIGURA(arq, legenda):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(AQUI, arq), width=Cm(12))
    P(legenda, 'Caption')


def shade(cell, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(s)


def TABELA(linhas, larguras, legenda):
    tb = d.add_table(rows=len(linhas), cols=len(linhas[0])); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB')
        bd.append(e)
    tb._tbl.tblPr.append(bd); tb.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(th)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in tb.rows:  # linha não se parte entre páginas
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i, row in enumerate(linhas):
        for j, txt in enumerate(row):
            c = tb.cell(i, j); c.width = Cm(larguras[j])
            p = c.paragraphs[0]; p.style = st['Lumis Tabela']
            if i == 0:
                p.paragraph_format.keep_with_next = True
                r = p.add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, '0F2D3A')
            else:
                runs(p, txt)
                if i == len(linhas) - 1:
                    p.paragraph_format.keep_with_next = True
                if i % 2 == 0:
                    shade(c, 'F1F6F4')
    P(legenda, 'Caption')


# ---------------- abertura ----------------
DESTAQUE('A expansão leva para novos mercados quatro riscos que a Lumis ainda não resolveu no Brasil: '
         'equidade, privacidade, transparência e governança. O crescimento só se sustenta com condições '
         'prévias e com uma métrica pública que mostre se o dano está caindo.')
B('**Equidade:** idosos de CEP D/E têm 31,8% de falso negativo, três vezes o melhor subgrupo, e a faixa D/E '
  'reclama oito vezes mais por paciente que a faixa A/B.')
B('**Privacidade:** nenhum dos seis contratos de dados passou por revisão jurídica, e a política para dados '
  'de saúde está em elaboração desde 2024.')
B('**Compromisso público:** a Lumis passa a publicar, a cada trimestre, o falso negativo por subgrupo.')
P('Analisamos a expansão como o Vetor Capital a propõe, com dois novos países e o Lumis Insight virando '
  'plataforma, e dois casos do backlog: a expansão para o México e o módulo de risco de crédito para bancos '
  '[fonte: Cap. 2, 1.1 e Quadro 16].')

# ---------------- 1 ----------------
H1('1. O que é material')
P('Cada tema recebeu duas notas: o dano que o negócio pode causar a pessoas e ao ambiente, e o risco '
  'financeiro que o tema traz para a Lumis. É material o tema com nota Alta em pelo menos um dos eixos.')
FIGURA('fig1_materialidade.png', 'Figura 1. Matriz de materialidade dupla. Fonte: Cap. 2, Anexo A; notas da equipe.')
TABELA([
    ['Tema', 'Evidência principal', 'Dano a pessoas', 'Risco à Lumis'],
    ['T1 Equidade no acesso', 'Falso negativo de 31,8% em 60+ de CEP D/E; a faixa D/E concentra 64% das reclamações [fonte: Cap. 2, Quadros 10 e 17]', 'Alto', 'Alto'],
    ['T2 Privacidade e dados de saúde', '35,4% dos dados de clientes com autorização fraca; política sem versão aprovada desde 2024 [fonte: Cap. 2, Quadros 7 e 17]', 'Alto', 'Alto'],
    ['T3 Transparência', '640 mil priorizações por mês com revisão humana de 2%; 94% divulgado contra 87,6% em campo [fonte: Cap. 2, Quadros 9, 11 e 12]', 'Alto', 'Alto'],
    ['T4 Governança da IA', 'Sem comitê de ética ou de risco; só o CTO libera versões; 2 de 4 incidentes detectados pelo cliente [fonte: Cap. 2, Quadros 12 e 14]', 'Alto', 'Alto'],
    ['T5 Clima e retenção', 'eNPS de −31 no time técnico; rotatividade de 27% no time de dados [fonte: Cap. 2, Quadro 15]', 'Médio', 'Alto'],
    ['T6 Diversidade', 'Mulheres são 12% e pessoas negras 6% da liderança [fonte: Cap. 2, Quadro 17]', 'Médio', 'Médio'],
    ['T7 Energia', '1.240 MWh por ano, 75% em inferência; 62% de fonte renovável [fonte: Cap. 2, Quadro 17]', 'Baixo', 'Baixo'],
], [3.3, 8.5, 2.1, 2.1], 'Tabela 1. Temas, evidência e notas. Fonte: Cap. 2, Anexo A; notas da equipe.')
P('Os temas T1 a T4 têm nota Alta nos dois eixos e já são problemas hoje. A expansão amplia cada um deles: '
  'mais clientes no mesmo modelo, populações sem histórico na base e dados de saúde cruzando fronteiras. '
  'Diversidade e energia ficam em monitoramento, porque os dados não mostram efeito relevante hoje; o consumo '
  'de energia, porém, cresce junto com a inferência [hipótese].')

# ---------------- 2 ----------------
H1('2. Quem pode ser prejudicado')
P('O erro cresce com a idade e com a faixa de CEP. No grupo 60+ de CEP D/E, quase um em cada três pacientes '
  'que precisavam de prioridade não a recebe, e a empresa nunca cruzou as reclamações com esse desempenho '
  '[fonte: Cap. 2, Quadros 10 e 17].')
P('A causa provável acompanha o modelo para qualquer mercado. Cinco variáveis que medem acesso ao sistema de '
  'saúde tanto quanto gravidade (atendimentos, custo acumulado, CEP, cobertura e faltas) somam 55,1% do peso '
  '[fonte: Cap. 2, Quadro 8; cálculo]. Quem usa menos o sistema parece menos grave e cai na fila [hipótese, '
  'coerente com o diagnóstico da Fase 1].')
TABELA([
    ['Frente', 'Grupo em risco', 'Por quê'],
    ['Plataforma e novos países', 'Pessoas de baixa renda e com pouco acesso à rede de saúde', 'As variáveis de acesso penalizam quem usa menos o sistema [hipótese]. Perfil de acesso no país de destino: [não consta]'],
    ['México', 'Subgrupos que a Lumis ainda não consegue identificar', 'Sem dado local de desfecho, a Lumis não sabe onde vai errar. O erro de hoje foi descoberto pelo cliente [fonte: Cap. 2, Quadro 14]'],
    ['Crédito para bancos', 'Moradores de CEP de baixa renda e pessoas com pouco histórico financeiro', 'CEP e cobertura pesam 16,3% no modelo de saúde; reaproveitados, podem negar crédito por endereço [hipótese]'],
], [3.4, 4.6, 8.0], 'Tabela 2. Grupos em risco em cada frente da expansão. Fonte: Cap. 2, Anexo A; leitura da equipe.')
P('O desempenho por sexo, raça ou cor e deficiência não é medido [não consta]. Não há como afirmar que esses '
  'grupos estão protegidos. Para medir o risco, propomos cinco indicadores:')
TABELA([
    ['Indicador', 'Limite proposto', 'Se passar do limite', 'Frequência'],
    ['E1. Falso negativo por subgrupo e razão entre o pior e o melhor', 'Até 1,5 vez (hoje: 3,0)', 'Revisão humana obrigatória no subgrupo; se persistir, suspensão daquela decisão, como prevê o C2', 'Quinzenal'],
    ['E2. Reclamações por 1.000 pacientes, por subgrupo, cruzadas com o E1', 'Nenhuma faixa acima de 2 vezes a média', 'Auditoria de equidade e investigação da causa', 'Mensal'],
    ['E3. Reversão humana por subgrupo (registro do C3)', 'Diferença de até 2 vezes entre subgrupos', 'Revisão das variáveis de acesso', 'Mensal'],
    ['E4. Representatividade dos dados no novo mercado', 'Amostra suficiente para medir o E1 em todo subgrupo', 'O mercado não entra em produção', 'Antes da entrada; depois, anual'],
    ['E5. Aprovação no crédito por faixa de CEP', 'Mínimo de 0,8 da taxa da faixa com mais aprovações¹', 'Revisão humana e retirada de CEP e cobertura do modelo', 'Mensal'],
], [5.0, 3.6, 5.2, 2.2], 'Tabela 3. Indicadores de equidade. Limites propostos pela equipe; E1 e E2 retomam os indicadores 1 e 5 da Entrega 2.')
P('A apuração fica com Yuri Nakamura (Dados), e a decisão de exigir revisão ou suspender o uso fica com o Head '
  'of AI Management, coordenador previsto no compromisso C4 [fonte: Cap. 2, Quadro 13; F1-E3]. Os cargos serão '
  'confirmados na Entrega 3.')

# ---------------- 3 ----------------
H1('3. Privacidade e dados sensíveis')
P('Dado de saúde é dado pessoal sensível na LGPD (art. 5º, II)², e quase toda a base da Lumis é desse tipo.')
B('**Revisão jurídica:** nenhum dos seis instrumentos de dados foi revisado desde a assinatura. Dos 5,9 mi de '
  'registros de clientes, 35,4% têm autorização fraca ou nenhuma, e a Prisma vence em 12/2026 [fonte: Cap. 2, Quadro 7].')
B('**Política:** a política para dados de saúde está em elaboração desde 2024, sem versão aprovada [fonte: Cap. 2, Quadro 17].')
B('**Perguntas sem resposta no caso:** os dados do Vila Ipê começam em 2019, três anos antes do contrato; a '
  'autorização da Sanare depende de anonimização irreversível (art. 12); a origem dos dados sintéticos [não consta].')
TABELA([
    ['LGPD', 'O que diz', 'Onde a Lumis fica exposta'],
    ['Art. 11', 'Dado sensível só com consentimento específico ou nas hipóteses listadas, como tutela da saúde por serviços de saúde', 'Não fica claro qual hipótese cobre treinar com dados de um hospital um produto vendido a outros [hipótese]'],
    ['Art. 11, § 4º', 'Vedado o uso compartilhado de dado de saúde para vantagem econômica, salvo na prestação de serviços de saúde', 'Plataforma que aprende com vários clientes e modelo de crédito treinado com dado de saúde [hipótese]'],
    ['Art. 11, § 5º', 'Operadoras de planos de saúde não podem usar dado de saúde para selecionar riscos', 'Classificação de sinistros para 9 seguradoras, 74 mil decisões por mês, se alguma for operadora de saúde [hipótese] [fonte: Cap. 2, Quadros 3 e 12]'],
    ['Art. 33', 'Transferência internacional só com proteção adequada no destino ou garantias, como cláusulas contratuais', 'Todo novo país. Lei aplicável no México: [não consta]'],
    ['Art. 52, II', 'Multa de até 2% do faturamento, limitada a R$ 50 mi por infração', 'O custo maior é retirar dados do treinamento e retreinar o modelo (Entrega 1)'],
], [2.4, 6.4, 7.2], 'Tabela 4. Pontos da LGPD que pesam na expansão. Fonte: Lei 13.709/2018; leitura da equipe, a confirmar com a DPO.')
P('Para o crédito, propomos uma regra fixa: dado de saúde não entra no modelo de crédito, que deve ser treinado '
  'só com dados financeiros autorizados. Hoje, só o Banco Meridiano autoriza esse uso [fonte: Cap. 2, Quadro 7].')
TABELA([
    ['Indicador', 'Hoje', 'Meta', 'Frequência'],
    ['P1. % dos registros de clientes em treinamento com autorização explícita e revisão jurídica', '0% com revisão jurídica; 64,6% com autorização explícita', '100% antes de qualquer frente nova; registro sem base sai do treinamento', 'Trimestral'],
    ['P2. % dos acessos a dados de pacientes revisados nos últimos 90 dias (C5)', '[não consta]', '100%', 'A cada 90 dias'],
    ['P3. Incidentes de privacidade e dias até a ação corretiva (C5)', 'Sem categoria própria no registro de incidentes [fonte: Cap. 2, Quadro 14]', 'Registro no dia em que é identificado; ação corretiva documentada', 'Mensal'],
], [5.6, 4.0, 4.4, 2.0], 'Tabela 5. Indicadores de privacidade. Responsável: Ana Beatriz Rangel (DPO), até a confirmação na Entrega 3.')

# ---------------- 4 ----------------
H1('4. A métrica que vamos publicar')
P('A Lumis não publica hoje nenhuma métrica de impacto [fonte: Cap. 2, Quadro 17]. Propomos publicar a taxa de '
  'falso negativo da priorização de atendimento, por subgrupo de idade e faixa de CEP. Foi esse o número que o '
  'Hospital Vila Ipê calculou por conta própria e enviou com a notificação, e ele mede o dano mais grave do produto: '
  'o paciente que precisava de prioridade e não a recebeu.')
TABELA([
    ['Item', 'Definição'],
    ['O que se publica', 'Falso negativo, sensibilidade e tamanho da amostra de cada subgrupo, mais a razão entre o pior e o melhor subgrupo e o valor do período anterior'],
    ['Ponto de partida', 'De 10,6% a 31,8%; razão de 3,0 vezes (1º semestre de 2026) [fonte: Cap. 2, Quadro 10]'],
    ['Meta', 'Razão de até 1,5 vez em 12 meses; nenhum subgrupo acima de 7,4%, o falso negativo da validação declarada, em 24 meses'],
    ['Periodicidade', 'Trimestral, no site da Lumis e no relatório de cada cliente; acompanhamento interno quinzenal, como prevê o C2'],
    ['Responsável', 'Head of AI Management publica; Yuri Nakamura (Dados) apura'],
    ['Verificação', 'Auditoria independente uma vez por ano'],
    ['Se piorar', 'Revisão humana obrigatória no subgrupo; suspensão do uso se persistir por dois ciclos quinzenais'],
], [3.4, 12.6], 'Tabela 6. Ficha da métrica pública. Metas propostas pela equipe.')
P('A sensibilidade é publicada junto para evitar que o número melhore às custas de marcar todos os pacientes '
  'como prioridade. Quando uma medida vira meta, ela deixa de ser boa medida (Goodhart) [fonte: Cap. 2, seção 2].')

# ---------------- 5 ----------------
H1('5. Condições para a expansão')
P('Recomendamos que nenhuma frente nova entre em produção antes de:')
N('falso negativo por subgrupo dentro do limite no Brasil, ou com revisão humana obrigatória nos subgrupos acima dele;')
N('validação local, por subgrupo, em cada novo mercado;')
N('política de privacidade aprovada e os seis instrumentos de dados revisados, com aditivos para a Prisma (antes de 12/2026) e o Vila Ipê (antes de 03/2027);')
N('relatório de impacto à proteção de dados para a plataforma e para cada país (LGPD, art. 38);')
N('no crédito, dado de saúde fora do modelo e diferença de aprovação por faixa de CEP acompanhada desde o primeiro dia;')
N('métrica pública de falso negativo por subgrupo já publicada.')
P('Nenhum ponto desta entrega contradiz os compromissos C1 a C5. A entrega dá limite e responsável aos '
  'compromissos C2 e C5, e a métrica pública cumpre o C1 com número. As condições acima seguem para o memorando '
  'ao conselho como condições prévias ao aporte.')

# ---------------- referências ----------------
P('Referências', 'Lumis Rótulo')
P('1. ESTADOS UNIDOS. 29 CFR § 1607.4: Information on impact (Uniform Guidelines on Employee Selection Procedures). '
  'Disponível em: https://www.law.cornell.edu/cfr/text/29/1607.4. Acesso em: 7 out. 2026. Usada por analogia: a regra '
  'foi criada para seleção de emprego.', 'Lumis Nota')
P('2. BRASIL. Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais (LGPD). Texto compilado. '
  'Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm. Acesso em: 7 out. 2026.', 'Lumis Nota')
P('Dados da Lumis: Cap. 2, Anexo A. Cálculos, notas e limites são da equipe. O detalhamento está no material de apoio '
  'da pasta da entrega. A IA generativa (Claude Code, Anthropic) apoiou a análise, a conferência das fontes e a '
  'redação, com revisão da equipe.', 'Lumis Nota')

d.save(OUT)
print('ok', OUT)
