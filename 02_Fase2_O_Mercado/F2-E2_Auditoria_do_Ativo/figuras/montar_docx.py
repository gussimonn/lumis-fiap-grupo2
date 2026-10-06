"""F2-E2 v1: transposição fiel da v1 do colega (PDF) para o padrão visual da F2-E1 v2.
O texto é o do colega, sem revisão. Base de formatação: F2-E1 v2 (capa só com título e subtítulo)."""
import re, sys, docx
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = 'C:/Users/gusta/Downloads/LumisOS/'
MODELO = BASE + '02_Fase2_O_Mercado/F2-E1_Mapa_do_Territorio/F2-E1_Mapa_do_Territorio_v2.docx'
FIG = BASE + '02_Fase2_O_Mercado/F2-E2_Auditoria_do_Ativo/figuras/'
OUT = sys.argv[1]

d = docx.Document(MODELO)
st = d.styles
body = d.element.body

# cabeçalho das páginas internas
runs_h = d.sections[0].header.paragraphs[0].runs
runs_h[1].text = '\tF2-E2 · Auditoria do Ativo'
runs_h[2].text = ''

# título e subtítulo; apaga o resto do corpo
d.paragraphs[0].runs[0].text = 'Auditoria do Ativo: Dados e Métricas'
for r in d.paragraphs[0].runs[1:]:
    r.text = ''
d.paragraphs[1].runs[0].text = 'Lumis Intelligence | Fase 2'
for r in d.paragraphs[1].runs[1:]:
    r.text = ''
sub = d.paragraphs[1]._p
apagar = False
for el in list(body):
    if apagar and el.tag != qn('w:sectPr'):
        body.remove(el)
    if el is sub:
        apagar = True

TAG = re.compile(r'(\[(?:fonte|hipótese|não consta|risco)[^\]]*\])')
TAGSTYLE = {'fonte': 'Tag Fonte', 'hipótese': 'Tag Hipótese', 'não consta': 'Tag Não consta', 'risco': 'Tag Risco'}


def runs(p, text):
    for i, seg in enumerate(text.split('**')):
        for part in TAG.split(seg):
            if not part:
                continue
            r = p.add_run(part)
            if TAG.fullmatch(part):
                r.style = st[TAGSTYLE[next(k for k in TAGSTYLE if part[1:].startswith(k))]]
            elif i % 2:
                r.bold = True
    return p


def P(text, style='Normal'):
    return runs(d.add_paragraph(style=style), text)


def H1(t): d.add_heading(t, level=1)


def DESTAQUE(rotulo, texto):
    p = d.add_paragraph(style='Lumis Destaque')
    r = p.add_run(rotulo + ' — '); r.bold = True
    p.add_run(texto)


def FIGURA(arq, legenda):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(FIG + arq, width=Cm(16))
    P(legenda, 'Caption')


def shade(cell, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(s)


def TABELA(linhas, larguras, legenda):
    tb = d.add_table(rows=len(linhas), cols=len(linhas[0])); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB')
        bd.append(e)
    tb._tbl.tblPr.append(bd); tb.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(th)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in tb.rows:  # linha não se parte entre páginas
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i, row in enumerate(linhas):
        for j, txt in enumerate(row):
            c = tb.cell(i, j); c.width = Cm(larguras[j])
            p = c.paragraphs[0]; p.style = st['Lumis Tabela']
            if i == 0:
                p.paragraph_format.keep_with_next = True  # cabeçalho não fica sozinho no pé da página
                r = p.add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, '0F2D3A')
            else:
                runs(p, txt)
                if i == len(linhas) - 1:
                    p.paragraph_format.keep_with_next = True  # legenda não fica sozinha na página seguinte
                if i % 2 == 0:
                    shade(c, 'F1F6F4')
    P(legenda, 'Caption')


# ---------------- corpo (texto da v1 do colega) ----------------
P('Esta entrega audita a composição do ativo de dados da Lumis, a robustez contratual de sua base, as principais variáveis do modelo, o desempenho real em produção, as métricas usadas comercialmente e um conjunto enxuto de indicadores de gestão capazes de alterar decisões.')

H1('1. Síntese executiva')
P('A auditoria aponta três fragilidades centrais. Primeiro, parte relevante dos dados de clientes utilizados no treinamento não possui autorização robusta: 2,09 milhões de registros têm base contratual frágil ou inexistente para treinamento, equivalentes a 35,4% dos dados de origem contratual e aproximadamente 9,5% da base total de 22 milhões de registros. Nenhum dos seis instrumentos informados passou por revisão jurídica desde a assinatura.')
P('Segundo, algumas variáveis de alto peso não medem diretamente a condição clínica que pretendem representar. Número de atendimentos, custo acumulado, tempo entre consulta e exame, faixa de CEP, tipo de cobertura e faltas podem funcionar como proxies de acesso, condição socioeconômica, organização da rede ou capacidade de pagamento. O risco é transformar desigualdades de acesso em sinais aparentemente clínicos.')
P('Terceiro, o desempenho observado em campo é materialmente inferior ao conjunto de validação de 2023: a acurácia caiu 6,5 pontos percentuais, a sensibilidade caiu 10,3 p.p. e o falso negativo subiu 10,3 p.p. O problema é ainda mais grave por subgrupo: pacientes com 60 anos ou mais em CEP D/E apresentam 31,8% de falso negativo, exatamente três vezes a taxa do melhor subgrupo, 18–59 anos em CEP A/B (10,6%).')
DESTAQUE('Conclusão executiva', 'o ativo de dados da Lumis é relevante em escala, mas não pode ser tratado como plenamente defensável sem corrigir fragilidades contratuais, monitorar proxies e substituir métricas comerciais frágeis por indicadores que governem decisões reais.')

H1('2. Origem dos dados e base contratual/legal')
P('A base de treinamento soma 22,0 milhões de registros. Desses, 5,9 milhões (26,8%) são de clientes; o restante vem de DATASUS e dados sintéticos internos. O ponto crítico não é apenas a quantidade, mas o direito de uso para treinamento e a atualidade da revisão jurídica.')
TABELA([
    ['Fonte', 'Volume', 'Situação para treinamento', 'Vencimento', 'Leitura de risco'],
    ['Hospital Vila Ipê', '1,2 mi atendimentos', 'Cláusula genérica: "melhoria contínua do serviço"', '03/2027', 'Frágil / ambígua'],
    ['Rede Sanare', '3,4 mi atendimentos', 'Sim, uso agregado e anonimizado', '08/2028', 'Mais robusta'],
    ['Seguradora Prisma', '890 mil sinistros', 'Contrato silente', '12/2026', 'Crítica / imediata'],
    ['Banco Meridiano', '410 mil operações', 'Sim, com auditoria anual do cliente', '05/2028', 'Mais robusta'],
    ['DATASUS', '14,0 mi registros', 'Uso público; treinamento autorizado no case', '—', 'Sem fragilidade contratual indicada'],
    ['Dados sintéticos internos', '2,1 mi registros', 'Geração própria; treinamento autorizado', '—', 'Controle interno'],
], [3.0, 2.9, 4.8, 2.3, 3.0], 'Tabela 1. Origem dos dados e situação para treinamento.')
DESTAQUE('Prioridade contratual', 'Seguradora Prisma deve ser tratada primeiro: o contrato é silente e vence em 12/2026. Hospital Vila Ipê também exige correção, pois a cláusula genérica não oferece a mesma clareza de uma autorização específica para treinamento. A renovação não deve ser o único gatilho: a revisão jurídica precisa ocorrer antes da continuidade do uso em treinamento.')

H1('3. Variáveis, representação e risco de proxy')
P('A pergunta central não é apenas "qual o peso da variável?", mas "o que ela realmente representa?". Variáveis administrativas e socioeconômicas podem carregar desigualdades históricas para dentro do modelo.')
TABELA([
    ['Variável', 'Peso', 'O que a empresa diz medir', 'O que pode medir de fato / distorção', 'Risco'],
    ['Nº atendimentos 24m', '18,4%', 'Histórico de necessidade de cuidado', 'Uso/acesso ao serviço; não informa gravidade de cada atendimento', 'Alto'],
    ['Custo acumulado', '15,1%', 'Gravidade do quadro clínico', 'Também reflete preços, cobertura e acesso', 'Alto'],
    ['Idade', '12,7%', 'Idade', 'Mede diretamente idade; risco está no tratamento discriminatório do efeito', 'Médio'],
    ['Nº comorbidades', '11,2%', 'Carga de doença', 'Pode depender de qualidade/volume do registro clínico', 'Médio'],
    ['Tempo consulta→exame', '9,8%', 'Urgência percebida', 'Pode refletir fila, capacidade da rede e acesso', 'Alto'],
    ['Faixa de CEP', '8,9%', 'Região de residência', 'Proxy potencial de renda, acesso, infraestrutura e perfil socioeconômico', 'Muito alto'],
    ['Tipo de plano/cobertura', '7,4%', 'Cobertura contratada', 'Pode funcionar como proxy de condição socioeconômica e acesso', 'Alto'],
    ['Painel laboratorial', '6,9%', 'Estado clínico atual', 'Mais diretamente clínico, mas depende de disponibilidade de exames', 'Médio'],
    ['Faltas em consultas', '5,3%', 'Adesão ao tratamento', 'Pode refletir transporte, trabalho, renda e barreiras de acesso', 'Alto'],
    ['Especialidade de origem', '4,3%', 'Via de entrada', 'Pode refletir organização e disponibilidade da rede', 'Médio'],
], [3.2, 1.4, 3.4, 6.2, 1.8], 'Tabela 2. Variáveis do modelo, o que medem e o risco de proxy.')
DESTAQUE('Proxy explícito', 'Faixa de CEP é o caso mais evidente: o modelo usa uma variável geográfica com peso de 8,9%, mas ela pode carregar renda, acesso a serviços e infraestrutura. Como o pior desempenho aparece justamente em CEP D/E, a variável deve ser auditada e testada quanto à contribuição para disparidades, sem assumir causalidade apenas pela correlação.')

H1('4. Validação declarada versus desempenho em campo')
TABELA([
    ['Recorte', 'Amostra', 'Acurácia', 'Sensibilidade', 'Falso negativo'],
    ['Validação 2023 — 2 hospitais, mesma região', '48.000', '94,1%', '92,6%', '7,4%'],
    ['Campo — 1º semestre de 2026', '1.940.000', '87,6%', '82,3%', '17,7%'],
    ['Variação', '—', '−6,5 p.p.', '−10,3 p.p.', '+10,3 p.p.'],
], [5.6, 2.6, 2.4, 2.8, 2.6], 'Tabela 3. Desempenho declarado e em campo.')
P('A diferença entre validação e produção mostra perda de desempenho fora do ambiente original. Para priorização clínica, sensibilidade e falso negativo são particularmente críticos porque indicam quantos pacientes que realmente deveriam ser priorizados estão sendo identificados — ou deixados para trás.')
FIGURA('fig1_subgrupo.png', 'Figura 1. Desempenho em campo por subgrupo — 1º semestre de 2026.')
P('Leitura do gráfico: o falso negativo cresce conforme se combinam maior idade e CEPs C/D/E. O grupo 60+ D/E chega a 31,8%, contra 10,6% em 18–59 A/B. O case informa que esse foi o dado identificado pelo próprio Hospital Vila Ipê e levado à notificação formal.')

H1('5. Revisão crítica das métricas apresentadas ao mercado')
TABELA([
    ['Claim atual', 'Como foi apurado', 'O que autoriza concluir', 'O que NÃO autoriza concluir', 'Decisão'],
    ['"Acurácia de 94%"', 'Validação 2023; 2 hospitais; 48 mil registros', 'Descreve aquele conjunto de validação', 'Não representa o desempenho atual em campo, que é 87,6%', 'Reformular e contextualizar'],
    ['"Redução de 30% no tempo de triagem"', 'Piloto de 6 semanas; 1 hospital; sem controle', 'Houve redução observada naquele piloto', 'Não permite atribuir causalidade nem generalizar para toda a base', 'Suspender como claim amplo'],
    ['"Mais de 5 milhões de vidas analisadas"', '5,2 mi registros processados com reprocessamentos', 'Há escala de processamento', 'Não equivale a 5,2 mi pessoas únicas', 'Substituir por "registros processados" ou deduplicar'],
    ['"NPS 72"', '9 respondentes indicados pelo comercial', 'Descreve a amostra respondente', 'Não sustenta satisfação representativa da base de clientes', 'Refazer pesquisa com amostra adequada'],
    ['"Disponibilidade 99,9%"', '99,92% da API', 'Descreve disponibilidade da interface', 'Não comprova disponibilidade ponta a ponta do serviço', 'Trocar por disponibilidade E2E'],
], [2.9, 3.2, 3.1, 3.7, 3.1], 'Tabela 4. Métricas apresentadas ao mercado e o que cada uma autoriza concluir.')

H1('6. Novo conjunto de indicadores de gestão')
P('Os indicadores abaixo foram escolhidos porque podem alterar uma decisão concreta. A aplicação segue as quatro perguntas do Capítulo 1: "Medida em quê?", "Medida por quem?", "Muda alguma decisão?" e "Esconde qual distribuição?".')
TABELA([
    ['Indicador', 'Medida em quê?', 'Medida por quem?', 'Muda alguma decisão?', 'Esconde qual distribuição?'],
    ['1. Taxa de falso negativo em produção por subgrupo', 'Pacientes que deveriam ser priorizados e foram classificados como baixa prioridade; mensal, por idade × CEP.', 'Cálculo automatizado pela área de Dados e validação amostral independente/cliente.', 'Aciona investigação, recalibração, restrição de uso ou suspensão quando houver deterioração relevante.', 'Sim: nunca reportar apenas agregado; abrir por idade, CEP e cliente.'],
    ['2. Drift de desempenho vs. baseline aprovado', 'Diferença de sensibilidade e falso negativo entre versão aprovada e produção, por versão e período.', 'Área de Dados, com revisão do Head de Gestão de IA.', 'Pode bloquear nova versão, exigir rollback ou revalidação antes de continuidade.', 'Sim: abrir por subgrupo e cliente; o agregado pode mascarar regressões locais.'],
    ['3. Cobertura contratual válida para treinamento', '% dos registros de clientes usados em treinamento com autorização explícita e revisão jurídica vigente.', 'DPO/Jurídico, reconciliado com inventário de dados.', 'Pode excluir fonte do treinamento, congelar uso, renegociar ou renovar cláusulas.', 'Sim: abrir por fonte/cliente e volume; média total pode esconder contrato crítico.'],
    ['4. Disponibilidade ponta a ponta do serviço', 'Tempo em que o fluxo completo — não só API — está disponível ao usuário final.', 'Operações/Tecnologia, com evidência de monitoramento E2E.', 'Direciona correção de gargalos, capacidade, SLA e comunicação de incidentes.', 'Sim: abrir por cliente, componente e severidade.'],
    ['5. Taxa de reclamações por 1.000 pacientes, por subgrupo', 'Reclamações formais normalizadas pelo volume de pacientes processados em cada faixa de CEP e idade.', 'CS registra; Gestão de IA cruza com performance e valida tendências.', 'Aciona auditoria de equidade, investigação de causas e revisão de processo/modelo.', 'Sim: obrigatório por CEP/idade; número absoluto sem denominador distorce a leitura.'],
], [2.9, 3.4, 3.4, 3.2, 3.1], 'Tabela 5. Indicadores de gestão e as quatro perguntas.')
DESTAQUE('Ajuste importante', '"Auditorias contratuais concluídas" não foi mantido como indicador final porque mede atividade, não resultado. O indicador mais útil é a cobertura efetiva dos registros usados em treinamento por autorização explícita e revisão jurídica vigente. Assim, a métrica muda uma decisão sobre usar ou retirar dados do treinamento.')

H1('7. Decisões recomendadas a partir da auditoria')
TABELA([
    ['Frente', 'Decisão recomendada', 'Por quê'],
    ['Contratos', 'Revisar imediatamente Prisma e Vila Ipê; inventariar os seis instrumentos e registrar base de uso por fonte.', 'Reduz risco de treinamento sem autorização clara.'],
    ['Modelo', 'Auditar contribuição das variáveis proxy, especialmente CEP, custo, utilização e faltas; testar impacto por subgrupo.', 'Evita reproduzir desigualdades como se fossem sinais clínicos.'],
    ['Validação', 'Adotar monitoramento contínuo em produção e critérios de rollback/suspensão associados a drift e falso negativo.', 'Fecha a distância entre validação e campo.'],
    ['Comercial', 'Interromper claims que extrapolam evidência e substituir por métricas reproduzíveis, contextualizadas e auditáveis.', 'Reduz risco reputacional e de due diligence.'],
    ['Gestão', 'Instituir os cinco indicadores acima com responsáveis e periodicidade definidos.', 'Transforma métricas em mecanismos de decisão.'],
], [2.6, 8.4, 5.0], 'Tabela 6. Decisões recomendadas por frente.')

H1('8. Evidências utilizadas')
P('Base principal: Anexo A do case Lumis Intelligence — quadros de "Base de dados e contratos", "Variáveis e pesos", "Desempenho declarado e em campo", "Desempenho em campo por subgrupo" e "Painel comercial".')
P('Critério metodológico das métricas: Capítulo 1 — "A IA e o Mercado", seção 2.3, que define as quatro perguntas de avaliação de qualquer métrica: medida em quê, medida por quem, muda alguma decisão e esconde qual distribuição.')

d.save(OUT)
print('ok')
