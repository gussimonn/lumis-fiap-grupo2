"""F2-E2 v2: relatório ao conselho, a partir do levantamento (F2-E2_Levantamento_v1.md) e das decisões D-026 e D-027.
Base de formatação: F2-E1 v2 (capa só com título e subtítulo, cabeçalho e rodapé do design system)."""
import re, sys, docx
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = 'C:/Users/gusta/Downloads/LumisOS/'
MODELO = BASE + '02_Fase2_O_Mercado/F2-E1_Mapa_do_Territorio/F2-E1_Mapa_do_Territorio_v2.docx'
FIG = BASE + '02_Fase2_O_Mercado/F2-E2_Auditoria_do_Ativo/figuras/'
OUT = sys.argv[1]

d = docx.Document(MODELO)
st = d.styles
body = d.element.body

# cabeçalho das páginas internas
hp = d.sections[0].header.paragraphs[0]
feito = False
for r in hp.runs:
    if r.text.strip():
        r.text = '' if feito else '\tF2-E2 · Auditoria do Ativo'
        feito = True

# título e subtítulo; apaga o resto do corpo
for p, t in ((d.paragraphs[0], 'Auditoria do Ativo'),
             (d.paragraphs[1], 'Relatório ao Conselho: de que é feito o ativo da Lumis e que números sobrevivem a uma auditoria independente')):
    p.runs[0].text = t
    for r in p.runs[1:]:
        r.text = ''
sub = d.paragraphs[1]._p
apagar = False
for el in list(body):
    if apagar and el.tag != qn('w:sectPr'):
        body.remove(el)
    if el is sub:
        apagar = True

TAG = re.compile(r'(\[(?:fonte|hipótese|não consta|risco)[^\]]*\])')
TAGSTYLE = {'fonte': 'Tag Fonte', 'hipótese': 'Tag Hipótese', 'não consta': 'Tag Não consta', 'risco': 'Tag Risco'}


def runs(p, text):
    for i, seg in enumerate(text.split('**')):
        for part in TAG.split(seg):
            if not part:
                continue
            r = p.add_run(part)
            if TAG.fullmatch(part):
                r.style = st[TAGSTYLE[next(k for k in TAGSTYLE if part[1:].startswith(k))]]
            elif i % 2:
                r.bold = True
    return p


def P(text, style='Normal'):
    return runs(d.add_paragraph(style=style), text)


def B(text):
    return P(text, 'List Bullet')


def H1(t): d.add_heading(t, level=1)
def H2(t): d.add_heading(t, level=2)


def FIGURA(arq, legenda):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(FIG + arq, width=Cm(16))
    P(legenda, 'Caption')


def shade(cell, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(s)


def TABELA(linhas, larguras, legenda, junta=False):
    tb = d.add_table(rows=len(linhas), cols=len(linhas[0])); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB')
        bd.append(e)
    tb._tbl.tblPr.append(bd); tb.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(th)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in tb.rows:  # linha não se parte entre páginas
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i, row in enumerate(linhas):
        for j, txt in enumerate(row):
            c = tb.cell(i, j); c.width = Cm(larguras[j])
            p = c.paragraphs[0]; p.style = st['Lumis Tabela']
            if i == 0:
                p.paragraph_format.keep_with_next = True  # cabeçalho não fica sozinho no pé da página
                r = p.add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, '0F2D3A')
            else:
                partes = txt.split('\n')
                runs(p, partes[0])
                for extra in partes[1:]:
                    runs(c.add_paragraph(style='Lumis Tabela'), extra)
                if junta or i == len(linhas) - 1:  # junta: a tabela inteira fica na mesma página
                    for q in c.paragraphs:
                        q.paragraph_format.keep_with_next = True  # legenda não fica sozinha na página seguinte
                if i % 2 == 0:
                    shade(c, 'F1F6F4')
    P(legenda, 'Caption')


# ---------------- abertura ----------------
# corpo enxuto (07/10/2026): números só onde sustentam a decisão; detalhe no anexo.
# Incorpora as edições da equipe no Word: pedido ao conselho mais curto, Sanare e Meridiano
# só como "Autoriza com condição." e base legal sem o marcador de hipótese.
P('Hoje o ativo da Lumis não passa numa auditoria independente. A parte da base que só a Lumis tem está sob contratos frágeis, e o modelo erra mais justamente com quem mais precisa de prioridade.', 'Lumis Destaque')
P('Um terço dos dados de clientes vem de contratos que não autorizam com clareza o treino do modelo, e nenhum contrato de dados foi revisado pelo jurídico. Em campo, de cada 10 idosos de CEP D/E que precisavam de prioridade, 3 ficaram para trás, e quem descobriu foi o hospital. Nenhuma das cinco afirmações que a Lumis divulga ao mercado resiste a uma auditoria do jeito que é divulgada.')
P('**Pedido ao conselho:** Restringir a recomendação automática para pacientes com 60 anos ou mais de CEP C e D/E, com o custo descrito na seção 5. E tirar do material comercial os números que não se sustentam.')

# ---------------- 1. dados ----------------
H1('1. De onde vêm os dados e se a Lumis pode usá-los')
P('Dois terços da base são do DATASUS, que é público e está aberto a qualquer concorrente. O que só a Lumis tem são os dados de clientes, pouco mais de um quarto da base, e é nessa parte que estão os problemas.')
TABELA([
    ['Fonte', 'O que o contrato diz sobre treino', 'Vence', 'Leitura da equipe'],
    ['**Seguradora Prisma**', 'Nada (contrato silente)', '12/2026', 'Sem autorização. Vence em menos de três meses'],
    ['**Hospital Vila Ipê**', 'Cláusula genérica: "melhoria contínua do serviço"', '03/2027', 'Autorização frágil. É também o hospital que notificou o viés'],
    ['**Rede Sanare**', 'Sim, para uso agregado e anonimizado', '08/2028', 'Autoriza com condição.'],
    ['**Banco Meridiano**', 'Sim, com auditoria anual do cliente', '05/2028', 'Autoriza com condição.'],
    ['**DATASUS e sintéticos**', 'Uso público e geração própria', 'Não se aplica', 'Autorizados. O DATASUS é aberto a todos'],
], [3.4, 5.0, 2.0, 5.6], 'Tabela 1. As fontes da base de treinamento, da mais frágil para a mais segura. Fonte: Cap. 2, Quadro 7; leitura da equipe.')
B('**Os contratos mais frágeis são os que vencem primeiro.** Prisma e Vila Ipê somam um terço dos dados de clientes. Nenhum dos seis contratos de dados passou por revisão jurídica, e o modelo atual já foi treinado com esses dados. O risco existe hoje, antes de qualquer vencimento.')
B('**Base legal.** Dado de saúde é dado pessoal sensível, e a LGPD pede finalidade específica e informada ao titular². Uma cláusula de "melhoria contínua do serviço" dificilmente cobre treinar um produto vendido a outros clientes. Quem trata dados em nome do cliente só pode usá-los para a finalidade definida por ele³. Ao treinar o próprio modelo, a Lumis passa a decidir uma finalidade própria e precisaria de base legal própria.')
B('**Dados anteriores à empresa.** A Lumis foi fundada em 2022, e parte dos dados de clientes começa em 2019 [fonte: Cap. 2, Quadros 3 e 7]. Sob que cláusula esse histórico chegou [não consta]. É a primeira pergunta que uma due diligence faria.')
P('O que é defensável hoje são os dados da Sanare e do Meridiano, e mesmo eles dependem de condições que ninguém verificou. A recomendação é ter o parecer da DPO sobre os seis contratos antes do próximo treino, renegociar Prisma e Vila Ipê antes do vencimento, incluindo o histórico, e aprovar a política de dados de saúde, parada desde 2024.')

# ---------------- 2. variáveis ----------------
H1('2. O que o modelo mede de fato')
P('As duas variáveis que mais pesam no modelo, número de atendimentos e custo acumulado, deveriam medir necessidade de cuidado e gravidade. Na prática, medem quem conseguiu ser atendido e quanto isso custou. Quem teve menos acesso parece menos grave para o modelo. O capítulo descreve esse caso e diz que "foi exatamente isso que aconteceu no incidente da Fase 1" [fonte: Cap. 2, 2.2].')
P('**Um caso real com o mesmo erro.** Em 2019, pesquisadores publicaram na revista Science a análise de um algoritmo comercial usado nos Estados Unidos para escolher pacientes para cuidado extra¹. O algoritmo previa quanto cada paciente ia custar, e não quão doente estava. Como pacientes negros tinham menos acesso a atendimento, gastavam menos e recebiam o mesmo escore que pacientes brancos mais saudáveis. Com a medida corrigida, a parcela de pacientes negros indicados para cuidado extra mais que dobraria. Na Lumis, o mecanismo é o mesmo, e quem fica para trás é quem tem menos acesso, como os idosos de CEP D/E. O caso não faz parte da história da Lumis e entra como comparação.')
TABELA([
    ['Variável', 'A empresa diz medir', 'O que tende a medir', 'Efeito'],
    ['Nº de atendimentos', 'Necessidade de cuidado', 'Quem conseguiu ser atendido', 'Quem tem menos acesso parece precisar de menos'],
    ['Custo acumulado', 'Gravidade', 'Acesso a procedimentos pagos', 'Quem gasta menos parece menos grave'],
    ['Faixa de CEP', 'Região de residência', 'Renda e oferta de serviços', 'Leva a desigualdade regional para o modelo'],
    ['Tipo de plano', 'Cobertura contratada', 'Renda', 'A prioridade passa a depender do plano'],
], [3.2, 3.6, 4.0, 5.2], 'Tabela 2. As variáveis que medem acesso. Juntas, são metade do peso do modelo. Fonte: Cap. 2, Quadro 8; leitura da equipe.')
P('O CEP reforça o caso: com a mesma idade, o erro cresce de A/B para D/E nas duas faixas etárias [fonte: Cap. 2, Quadro 10]. Por que os idosos são os mais afetados o anexo não permite separar, e os testes ficam para a Fase 3.')

# ---------------- 3. campo ----------------
H1('3. O que acontece em campo')
P('Os 94% de acurácia vêm de uma amostra de 2023 com dois hospitais da mesma região. Em campo, o número que importa para priorização, o paciente que precisava de prioridade e não recebeu, mais que dobrou [fonte: Cap. 2, Quadro 9].')
FIGURA('fig1_subgrupo_v2.png', 'Figura 1. Pacientes que precisavam de prioridade e não receberam, por grupo, 1º semestre de 2026. Fonte: Cap. 2, Quadros 9 e 10.')
P('O erro se concentra em quem já é mais vulnerável: o modelo deixa de priorizar idosos de CEP D/E três vezes mais que adultos de CEP A/B. Foi o Hospital Vila Ipê que mediu isso, e não a Lumis. Só 2% das decisões passam por revisão humana, e as reclamações dos pacientes de CEP D/E, que são a maioria, nunca foram cruzadas com o desempenho do modelo [fonte: Cap. 2, Quadros 10, 12 e 17].')

# ---------------- 4. métricas ----------------
H1('4. Os números que mostramos ao mercado')
P('O Vetor Capital avisou que número que não sobrevive a uma auditoria independente vira passivo. Na maior parte dos casos, o número está certo e a forma de divulgar é que não se sustenta.')
TABELA([
    ['Afirmação', 'O que de fato mostra', 'Destino'],
    ['"Acurácia de 94%"', 'Uma amostra de 2023 de dois hospitais, sem os grupos', 'Trocar pelo erro em campo por grupo (indicador 1)'],
    ['"Redução de 30% no tempo de triagem"', 'Um piloto de seis semanas, sem grupo de comparação', 'Retirar até haver estudo com comparação'],
    ['"Mais de 5 milhões de vidas analisadas"', 'Processamentos, com o mesmo paciente contado mais de uma vez', 'Retirar já'],
    ['"NPS 72"', '9 clientes indicados pelo comercial', 'Trocar pelas reclamações por grupo (indicador 3)'],
    ['"Disponibilidade de 99,9%"', 'Só a interface técnica, sem o serviço completo', 'Tirar do material comercial'],
], [4.4, 6.4, 5.2], 'Tabela 3. As cinco afirmações do painel comercial. As quatro perguntas aplicadas a cada uma estão no Anexo. Fonte: Cap. 2, Quadro 11.')
P('Os "5 milhões de vidas" são o caso mais grave: a duplicidade foi achada em janeiro, corrigida no sistema e continua no material comercial [fonte: Cap. 2, Quadro 14]. Manter as cinco afirmações como estão contradiz nosso compromisso de Transparência (C1).')

# ---------------- 5. indicadores ----------------
H1('5. O que passamos a medir e o que decidimos agora')
P('No lugar dos números de vitrine, propomos cinco indicadores. Cada um tem um ponto de ação, uma decisão que ele dispara e um responsável. Os pontos de ação são proposta da equipe, e os cargos são provisórios até a Entrega 3.')
TABELA([
    ['Indicador', 'Quando age [hipótese]', 'O que decide', 'Quem responde'],
    ['**1. Erro em campo por grupo** (idade, CEP e cliente)', 'Grupo com o dobro do erro do melhor grupo', 'Suspender a recomendação automática no grupo; barrar versão que piore algum grupo', 'Yuri Nakamura mede; a CEO decide'],
    ['**2. Direito de uso dos dados** (por contrato)', 'Dado sem autorização e parecer jurídico; contrato a seis meses do fim', 'O que entra no próximo treino; ordem de renegociação', 'Ana Beatriz Rangel (DPO)'],
    ['**3. Reclamações por faixa de CEP**', 'Grupo com uma vez e meia mais reclamações do que o seu peso na base', 'Revisão clínica dos casos e análise do grupo no indicador 1', 'Head of AI Management'],
    ['**4. Incidentes descobertos antes do cliente**', 'Cliente descobre primeiro, ou incidente aberto há mais de 30 dias', 'Rever o monitoramento; levar o caso ao conselho', 'Head of AI Management'],
    ['**5. Afirmações públicas auditáveis**', 'Afirmação que alguém de fora não consegue refazer', 'Retirar do material antes da próxima apresentação', 'Camila Torres propõe; Head of AI aprova'],
], [4.2, 4.0, 4.6, 3.2], 'Tabela 4. Os cinco indicadores. Valores de hoje, as quatro perguntas e os indicadores descartados estão no Anexo. Fonte: Cap. 2, Quadros 7, 10, 13, 14 e 17; propostas da equipe.')
P('Como a Lumis mede o erro do próprio produto, cada indicador tem conferência de fora: o cliente refaz a conta do próprio grupo, como o Vila Ipê já fez, e um auditor confere a cada trimestre. Para a revisão quinzenal ter casos suficientes, a revisão humana sobe de 2% para 10% nos grupos críticos. O NIST pede que todo sistema de IA tenha responsável e meio definidos para ser desligado⁴.')
P('**A decisão de agora.** Dois grupos já passam do ponto de ação do indicador 1: idosos de CEP C e de CEP D/E. Pelo nosso compromisso de Não Amplificação de Danos (C2), o uso é suspenso até a correção quando há padrão de viés. Por isso recomendamos restringir já a recomendação automática nesses grupos (D-026). Cerca de um quarto das decisões de priorização volta à triagem do próprio hospital, enquanto o modelo segue rodando em paralelo, sem decidir, para medir a correção. Quanto isso pesa em cada hospital [não consta], e a Entrega 3 define quem executa a restrição.')

# ---------------- 6. tese ----------------
H1('6. O que isso muda na tese')
P('A base de clientes só vira vantagem difícil de copiar quando tiver direito de uso limpo e erro por grupo que alguém de fora consiga refazer, como a Entrega 1 já apontava. Por isso, o memorando leva três condições prévias ao aporte: regularizar os contratos antes do vencimento, tirar do mercado os números que não se sustentam e medir o erro por grupo com conferência externa.')

H2('Referências')
for t in [
    '1. OBERMEYER, Z. et al. Dissecting racial bias in an algorithm used to manage the health of populations. Science, v. 366, n. 6464, p. 447-453, 2019.',
    '2. BRASIL. Lei nº 13.709/2018 (LGPD), arts. 5º, 6º, 11 e 52. Consulta em 06/10/2026.',
    '3. ANPD. Guia Orientativo para Definições dos Agentes de Tratamento de Dados Pessoais e do Encarregado. Versão 2.0, 2022.',
    '4. NIST. Artificial Intelligence Risk Management Framework (AI RMF 1.0). NIST AI 100-1, 2023.',
]:
    P(t, 'Lumis Nota')
P('Dados da Lumis: Cap. 2, Anexo A. Contas e propostas são da equipe. As empresas reais citadas são análogos e não fazem parte do caso Lumis.', 'Lumis Nota')

# ---------------- anexo ----------------
H1('Anexo: as quatro perguntas')
d.paragraphs[-1].paragraph_format.page_break_before = True  # sem parágrafo vazio, que gerava página em branco

H2('A1. As quatro perguntas aplicadas às afirmações do mercado')
TABELA([
    ['Afirmação', 'Medida em quê?', 'Medida por quem?', 'Muda alguma decisão?', 'Esconde qual distribuição?'],
    ['Acurácia de 94%', '48 mil registros de 2023, dois hospitais da mesma região', 'Quem apurou [não consta]', 'Não. Virou discurso comercial', 'O erro em campo e os grupos: 31,8% no pior'],
    ['Redução de 30% na triagem', 'Seis semanas, um hospital, sem comparação', '[não consta]', 'Não, sem comparação', 'Não abre por grupo nem por hospital'],
    ['5 milhões de vidas', 'Registros processados, com repetição', 'Duplicidade achada pela auditoria interna em 01/2026', 'Não. É volume', 'Quantas pessoas únicas há [não consta]'],
    ['NPS 72', '9 respondentes', 'Escolhidos pelo comercial', 'Não', 'As 38 contas e as 74 reclamações'],
    ['Disponibilidade de 99,9%', 'Só a interface técnica', '[não consta]', 'Pouco. É operacional', 'O serviço completo e a qualidade da resposta'],
], [2.8, 3.4, 3.2, 3.0, 3.6], 'Tabela A1. Fonte: Cap. 2, Quadros 9, 11, 14 e 17; as quatro perguntas: Cap. 2, 2.3.')

H2('A2. As quatro perguntas aplicadas aos indicadores propostos')
TABELA([
    ['Indicador e valor de hoje', 'Medida em quê?', 'Medida por quem?', 'Muda alguma decisão?', 'Esconde qual distribuição?'],
    ['1. Erro por grupo: pior grupo 3,0 vezes o melhor', 'Revisão humana (10% nos grupos críticos) e desfecho do cliente', 'Pela Lumis. Cliente e auditor refazem', 'Sim: restrição de grupo e liberação de versão', 'Diferenças por cliente; grupos que o caso não abre, como sexo'],
    ['2. Direito de uso: 0% com parecer jurídico', 'Contratos contra a base de treino', 'Pela DPO. Parecer externo confere', 'Sim: treino e renegociação', 'A concentração na Sanare e o período antes do contrato'],
    ['3. Reclamações: CEP D/E com 2,5 vezes o seu peso', 'Registro formal, janela de 90 dias', 'Time de sucesso do cliente, cargo [não consta]. O cliente confere', 'Sim: revisão clínica', 'Quem não reclama'],
    ['4. Incidentes: 2 de 4 descobertos pelo cliente', 'Registro de incidentes e notificações dos clientes', 'Head of AI Management. A DPO confere', 'Sim: monitoramento e liberação', 'A gravidade e o incidente que ninguém viu'],
    ['5. Afirmações auditáveis: 0 de 5', 'Material público contra checklist de relato', 'Por quem não vende. Cliente ou auditor refaz', 'Sim: publicar ou retirar', 'O grupo que quem escolhe o recorte deixa de fora'],
], [3.2, 3.2, 3.4, 2.8, 3.4], 'Tabela A2. Fonte: análise da equipe sobre o Cap. 2, Quadros 7, 10, 13, 14 e 17.')

H2('A3. As quatro perguntas aplicadas aos indicadores descartados')
TABELA([
    ['Candidato', 'Medida em quê?', 'Medida por quem?', 'Muda alguma decisão?', 'Esconde qual distribuição?'],
    ['Acurácia global em campo', 'Base completa', '[não consta]', 'Pouco', 'O erro grave e os grupos'],
    ['Comparação com a versão aprovada', 'Troca de versão', 'Pelo CTO, que também libera a versão', 'Repete o indicador 1 e vira a regra de liberação dele', 'O mesmo que o indicador 1'],
    ['Disponibilidade de ponta a ponta', 'Cadeia técnica', 'Pelo CTO', 'Decisão técnica, que não chega ao conselho', 'Não abre por grupo'],
    ['Auditorias contratuais concluídas', 'Número de auditorias', 'Pela DPO', 'Não. Mede atividade', 'Quantos registros estão afetados'],
    ['Pacientes únicos analisados', 'Base sem duplicidade', 'Pela área de Dados', 'Não. É volume', 'Quem foi mal atendido'],
], [3.4, 2.8, 3.2, 3.6, 3.0], 'Tabela A3. Fonte: análise da equipe sobre o Cap. 2, Quadros 9, 11 e 13.', junta=True)

d.save(OUT)
print('ok')
