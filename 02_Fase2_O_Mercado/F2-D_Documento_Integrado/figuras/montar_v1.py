"""F2-D v1: monta o documento único da Fase 2 (dossiê ao conselho e ao comitê do Vetor Capital) e o memorando
ao conselho (F2-M), no padrão visual da Lumis (00_Lumis/Design_System/Modelo_Entrega_Lumis.docx).

Uso:
    python montar_v1.py dossie <saida.docx>   # documento único
    python montar_v1.py memo   <saida.docx>   # memorando ao conselho
    python montar_v1.py lint                  # confere o texto (códigos internos, travessões, chaves)

O texto de cada parte está no fim deste arquivo (CONTEUDO_DOSSIE e CONTEUDO_MEMO), uma instrução por linha:
# título 1, ## título 2, > destaque, - item, 1. item numerado, FIG[chave]: legenda, TAB[chave]: larguras | legenda
seguida das linhas "| ... |", NIVEIS[...] (tabela por níveis), CAIXA ... FIMCAIXA (prompt), QUEBRA, REFERENCIAS,
NOTA:, LINHA:, TITULO:, SUBTITULO:. Marcadores [fonte: ...], [hipótese], [não consta] e [risco: ...] saem com os
estilos Tag do modelo; {fig:chave}, {tab:chave} e {ref:chave} viram números.
As figuras são as das entregas, menos as quatro da Entrega 2, refeitas para o dossiê por gerar_figuras_e2_dossie.py.
Decisões: D-042 a D-044.
"""

import os, re, sys
import docx
from docx.shared import Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as VA
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = 'C:/Users/gusta/Downloads/LumisOS/'
MODELO = BASE + '00_Lumis/Design_System/Modelo_Entrega_Lumis.docx'
F2 = '02_Fase2_O_Mercado/'
LARGURA_UTIL = 16.0  # cm (A4 com margens de 2,5 cm)

META = {
    'dossie': {
        'rotulo': 'Fase 2 · O Mercado',
        'titulo': 'O mapa e a direção',
        'subtitulo': 'Dossiê ao conselho e ao comitê de investimento do Vetor Capital: onde a Lumis deve crescer, com que ativo e sob que condições',
        'equipe': 'Equipe: Bruno Müller, Diego Franca Evangelista, Felipe Alef, Gustavo Halfen Simon e Maria Fernanda Barros',
        'data': 'Data: 08/10/2026',
        'cabecalho': 'Fase 2 · O mapa e a direção',
    },
    'memo': {
        'cabecalho': 'Fase 2 · Memorando ao conselho',
    },
}

# Figuras disponíveis (as mesmas imagens que estão nas entregas vigentes). Largura padrão em cm.
FIGS = {
    'e1_camadas':      (F2 + 'F2-E1_Mapa_do_Territorio/figuras/fig1_camadas.png', 12.0),
    'e1_ameacas':      (F2 + 'F2-E1_Mapa_do_Territorio/figuras/fig2_ameacas.png', 10.5),
    'e1_margem':       (F2 + 'F2-E1_Mapa_do_Territorio/figuras/fig4_margem.png', 14.0),
    'e2_base':         (F2 + 'F2-D_Documento_Integrado/figuras/d_e2_base.png', 13.0),
    'e2_proxy':        (F2 + 'F2-D_Documento_Integrado/figuras/d_e2_proxy.png', 13.0),
    'e2_subgrupo':     (F2 + 'F2-E2_Auditoria_do_Ativo/figuras/fig1_subgrupo_v2.png', 13.0),
    'e2_metricas':     (F2 + 'F2-D_Documento_Integrado/figuras/d_e2_metricas.png', 13.0),
    'e2_indicadores':  (F2 + 'F2-D_Documento_Integrado/figuras/d_e2_indicadores.png', 13.0),
    'e4_clima':        (F2 + 'F2-E4_Cultura_e_Funil_de_Inovacao/figuras/v1_fig1_clima.png', 14.0),
    'e4_funil':        (F2 + 'F2-E4_Cultura_e_Funil_de_Inovacao/figuras/v1_fig2_funil.png', 15.0),
    'e4_capacidade':   (F2 + 'F2-E4_Cultura_e_Funil_de_Inovacao/figuras/v1_fig3_capacidade.png', 14.0),
    'e5_materialidade': (F2 + 'F2-E5_Due_Diligence_ESG/figuras/fig1_materialidade.png', 11.0),
}

# Referências externas. Entram numeradas na ordem da primeira citação {ref:chave}.
REFS = {
    'shapiro': 'SHAPIRO, C.; VARIAN, H. R. Information Rules: A Strategic Guide to the Network Economy. Boston: Harvard Business School Press, 1999.',
    'mv_klas': 'MV. MV se torna a 5ª maior fornecedora global de prontuário eletrônico hospitalar (dados KLAS 2025). 4 ago. 2025. Disponível em: https://mv.com.br/imprensa/mv-se-torna-a-5a-maior-fornecedora-global-de-prontuario-eletronico-hospitalar. Acesso em: 5 out. 2026.',
    'einstein': 'CONVERGÊNCIA DIGITAL. Hospital Israelita Albert Einstein: algoritmos, IA e inovação salvam vidas. 8 jan. 2025. Disponível em: https://convergenciadigital.com.br/mercado/hospital-israelita-albert-einstein-algoritmos-ia-e-inovacao-salvam-vidas/. Acesso em: 5 out. 2026.',
    'wong': 'WONG, A. et al. External Validation of a Widely Implemented Proprietary Sepsis Prediction Model in Hospitalized Patients. JAMA Internal Medicine, v. 181, n. 8, p. 1065-1070, 2021.',
    'techcrunch': 'TECHCRUNCH. Anthropic hikes the price of its Haiku model. 4 nov. 2024. Disponível em: https://techcrunch.com/2024/11/04/anthropic-hikes-the-price-of-its-haiku-model. Acesso em: 5 out. 2026.',
    'obermeyer': 'OBERMEYER, Z. et al. Dissecting racial bias in an algorithm used to manage the health of populations. Science, v. 366, n. 6464, p. 447-453, 2019.',
    'lgpd': 'BRASIL. Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais (LGPD). Texto compilado. Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm. Acesso em: 7 out. 2026.',
    'anpd': 'ANPD. Guia Orientativo para Definições dos Agentes de Tratamento de Dados Pessoais e do Encarregado. Versão 2.0, 2022.',
    'cfr': 'ESTADOS UNIDOS. 29 CFR § 1607.4: Information on impact (Uniform Guidelines on Employee Selection Procedures). Disponível em: https://www.law.cornell.edu/cfr/text/29/1607.4. Acesso em: 7 out. 2026. Usada por analogia: a regra foi criada para seleção de emprego.',
    'schein': 'SCHEIN, E. H.; SCHEIN, P. A. Organizational Culture and Leadership. 5. ed. Hoboken: Wiley, 2017. A bibliografia do curso registra a edição como de 2016.',
    'cooper': 'COOPER, R. G. Stage-gate systems: a new tool for managing new products. Business Horizons, v. 33, n. 3, p. 44-54, 1990.',
    'nagji': 'NAGJI, B.; TUFF, G. Managing Your Innovation Portfolio. Harvard Business Review, maio 2012. Disponível em: https://hbr.org/2012/05/managing-your-innovation-portfolio.',
}

# Cores dos níveis da tabela "difícil de copiar" (mesmas da Entrega 1).
NIVEIS_COR = {
    'Copiável': ('6B7C85', None),
    'Parcialmente difícil': ('0F2D3A', 'F1F6F4'),
    'Moderadamente difícil': ('0F2D3A', None),
    'Valioso, mas travado': ('B7791F', 'FBF3E6'),
    'Difícil de copiar': ('1E8C6B', 'EEF7F3'),
}

TAG = re.compile(r'(\[(?:fonte|hipótese|não consta|risco)[^\]]*\])')
TAGSTYLE = {'fonte': 'Tag Fonte', 'hipótese': 'Tag Hipótese', 'não consta': 'Tag Não consta', 'risco': 'Tag Risco'}
SUP = re.compile(r'\x01(\d+(?:,\d+)*)\x02')
SUPMAP = str.maketrans('0123456789,', '⁰¹²³⁴⁵⁶⁷⁸⁹˒')
DIRETIVA = re.compile(r'^(FIG|TAB|NIVEIS)\[([^\]|]+)(?:\|([\d.,]+))?\]:\s*(.*)$')


# ------------------------------------------------------------------ leitura
def _ler_arquivos(caminho):
    if os.path.isdir(caminho):
        arqs = sorted(f for f in os.listdir(caminho) if f.endswith('.txt'))
        linhas = []
        for f in arqs:
            with open(os.path.join(caminho, f), encoding='utf-8') as fh:
                linhas += [(f, i + 1, l.rstrip('\n')) for i, l in enumerate(fh)]
        return linhas
    with open(caminho, encoding='utf-8') as fh:
        return [(os.path.basename(caminho), i + 1, l.rstrip('\n')) for i, l in enumerate(fh)]


def blocos(linhas):
    """Agrupa as linhas em blocos (tipo, dados, origem)."""
    out, i = [], 0
    while i < len(linhas):
        arq, n, l = linhas[i]
        s = l.strip()
        org = f'{arq}:{n}'
        if not s or s.startswith('%'):
            i += 1; continue
        m = DIRETIVA.match(s)
        if m:
            tipo, chave, larg, resto = m.group(1), m.group(2).strip(), m.group(3), m.group(4)
            if tipo == 'FIG':
                out.append(('FIG', dict(chave=chave, largura=float(larg.replace(',', '.')) if larg else None, legenda=resto.strip()), org))
                i += 1; continue
            partes = resto.split('|', 1)
            larguras = [float(x) for x in partes[0].replace(' ', '').split(';' if ';' in partes[0] else ',') if x] if len(partes) == 2 else []
            legenda = partes[1].strip() if len(partes) == 2 else resto.strip()
            linhas_tab = []
            i += 1
            while i < len(linhas) and linhas[i][2].strip().startswith('|'):
                cel = linhas[i][2].strip()
                cel = cel[1:-1] if cel.endswith('|') else cel[1:]
                linhas_tab.append([c.strip() for c in cel.split('|')])
                i += 1
            out.append((tipo, dict(chave=chave, larguras=larguras, legenda=legenda, linhas=linhas_tab), org))
            continue
        if s == 'CAIXA':
            conteudo = []
            i += 1
            while i < len(linhas) and linhas[i][2].strip() != 'FIMCAIXA':
                conteudo.append(linhas[i][2])
                i += 1
            i += 1
            out.append(('CAIXA', dict(linhas=conteudo), org)); continue
        out.append(('LINHA', dict(texto=s), org))
        i += 1
    return out


def ler_conteudo(caminho):
    if caminho == ':dossie':
        linhas = []
        for f in sorted(CONTEUDO_DOSSIE):
            linhas += [(f, i + 1, l) for i, l in enumerate(CONTEUDO_DOSSIE[f].split('\n'))]
        return linhas
    if caminho == ':memo':
        return [('memo.txt', i + 1, l) for i, l in enumerate(CONTEUDO_MEMO.split('\n'))]
    return _ler_arquivos(caminho)


# ------------------------------------------------------------------ numeração
def numerar(bs):
    figs, tabs, refs = {}, {}, {}
    for tipo, d, org in bs:
        if tipo == 'FIG':
            figs.setdefault(d['chave'], len(figs) + 1)
        elif tipo in ('TAB', 'NIVEIS'):
            tabs.setdefault(d['chave'], len(tabs) + 1)
    textos = []
    for tipo, d, org in bs:
        if tipo == 'LINHA':
            textos.append(d['texto'])
        elif tipo in ('TAB', 'NIVEIS'):
            textos += [c for row in d['linhas'] for c in row] + [d['legenda']]
        elif tipo == 'FIG':
            textos.append(d['legenda'])
    for t in textos:
        for k in re.findall(r'\{ref:([^}]+)\}', t):
            for kk in k.split(','):
                refs.setdefault(kk.strip(), len(refs) + 1)
    return figs, tabs, refs


def resolver(texto, figs, tabs, refs, avisos, org):
    def fig(m):
        k = m.group(1)
        if k not in figs:
            avisos.append(f'{org}: figura citada sem bloco FIG: {k}')
            return '??'
        return str(figs[k])

    def tab(m):
        k = m.group(1)
        if k not in tabs:
            avisos.append(f'{org}: tabela citada sem bloco TAB: {k}')
            return '??'
        return str(tabs[k])

    def ref(m):
        nums = []
        for k in m.group(1).split(','):
            k = k.strip()
            if k not in REFS:
                avisos.append(f'{org}: referência fora do registro: {k}')
            nums.append(str(refs.get(k, 0)))
        return '\x01' + ','.join(nums) + '\x02'

    texto = re.sub(r'\{fig:([^}]+)\}', fig, texto)
    texto = re.sub(r'\{tab:([^}]+)\}', tab, texto)
    texto = re.sub(r'\{ref:([^}]+)\}', ref, texto)
    if re.search(r'\{[a-z]+:', texto):
        avisos.append(f'{org}: marcador não resolvido: {texto[:80]}')
    return texto


# ------------------------------------------------------------------ verificação de texto
INTERNOS = [
    (re.compile(r'\bD-\d{2,3}\b'), 'código de decisão interna (D-xx)'),
    (re.compile(r'\bC[1-5]\b'), 'código de compromisso (C1 a C5)'),
    (re.compile(r'\bF\d-[EAMD]\d?\b'), 'código de entrega (FX-EY)'),
    (re.compile(r'\bQ-\d+\b'), 'código de pergunta interna (Q-x)'),
    (re.compile(r'DECISOES|ESTADO\.md|[Ll]evantamento_v\d|\.md\b'), 'nome de arquivo interno'),
    (re.compile(r'[—–]'), 'travessão (trocar por vírgula, dois-pontos, ponto ou parênteses)'),
    (re.compile(r'\b(crucial|robust[oa]s?|alavanc\w+|potencializ\w+|jornada|ressalt\w+|destaca\w*)\b', re.I), 'palavra marcada pelo humanizer'),
    (re.compile(r'não (?:é|são) [^.;:]{1,60}, e sim\b|\bnão só\b[^.]*\bmas também\b', re.I), 'contraste "não X, e sim Y"'),
    (re.compile(r'\{(?!fig:|tab:|ref:)'), 'chave estranha'),
]


def checar_texto(texto, org, avisos, celula=False):
    for rx, nome in INTERNOS:
        if celula and nome.startswith('travessão') and texto.strip() in ('—', '–'):
            continue
        m = rx.search(texto)
        if m:
            avisos.append(f'{org}: {nome}: "...{texto[max(0, m.start() - 30):m.end() + 30]}..."')
    if texto.count('**') % 2:
        avisos.append(f'{org}: negrito sem fechamento: {texto[:60]}')
    if texto.count('[') != texto.count(']'):
        avisos.append(f'{org}: colchete sem par: {texto[:60]}')


# ------------------------------------------------------------------ escrita no Word
class Doc:
    def __init__(self, modo):
        self.d = docx.Document(MODELO)
        self.st = self.d.styles
        self.body = self.d.element.body
        self.modo = modo
        self.lista_num_aberta = False
        self._preparar(modo)

    # capa, cabeçalho e limpeza do exemplo do modelo
    def _preparar(self, modo):
        d, meta = self.d, META[modo]
        hp = d.sections[0].header.paragraphs[0]
        feito = False
        for r in hp.runs:
            if r.text.strip():
                r.text = '' if feito else '\t' + meta['cabecalho']
                feito = True
        paras = d.paragraphs
        corte = None
        for p in paras:  # o parágrafo com a quebra de página fecha a capa
            if p._p.xpath('.//w:br[@w:type="page"]'):
                corte = p._p
                break
        if modo == 'dossie':
            def trocar(p, t):
                if p.runs:
                    p.runs[0].text = t
                    for r in p.runs[1:]:
                        r.text = ''
                else:
                    p.add_run(t)
            for p in paras:
                nome = p.style.name
                if p._p is corte:
                    break
                if nome == 'Lumis Rótulo':
                    trocar(p, meta['rotulo'])
                elif nome == 'Title':
                    trocar(p, meta['titulo'])
                elif nome == 'Subtitle':
                    trocar(p, meta['subtitulo'])
                elif nome == 'Lumis Nota' and p.text.startswith('Equipe'):
                    trocar(p, meta['equipe'])
                elif nome == 'Lumis Nota' and p.text.startswith('Data'):
                    trocar(p, meta['data'])
            apagar = False
            for el in list(self.body):
                if apagar and el.tag != qn('w:sectPr'):
                    self.body.remove(el)
                if el is corte:
                    apagar = True
        else:  # memorando: sem capa; cabeçalho e rodapé já na primeira página
            for el in list(self.body):
                if el.tag != qn('w:sectPr'):
                    self.body.remove(el)
            d.sections[0].different_first_page_header_footer = False

    # texto com negrito, marcadores de origem e números sobrescritos
    def runs(self, p, texto):
        for i, seg in enumerate(texto.split('**')):
            for parte in TAG.split(seg):
                if not parte:
                    continue
                if TAG.fullmatch(parte):
                    r = p.add_run(parte)
                    r.style = self.st[TAGSTYLE[next(k for k in TAGSTYLE if parte[1:].startswith(k))]]
                    continue
                pos = 0
                for m in SUP.finditer(parte):
                    if m.start() > pos:
                        r = p.add_run(parte[pos:m.start()])
                        if i % 2:
                            r.bold = True
                    r = p.add_run(m.group(1).translate(SUPMAP))
                    pos = m.end()
                if pos < len(parte):
                    r = p.add_run(parte[pos:])
                    if i % 2:
                        r.bold = True
        return p

    def par(self, texto, estilo='Normal'):
        p = self._aplicar_quebra(self.d.add_paragraph(style=estilo))
        return self.runs(p, texto)

    def numerado(self, texto, inicio=1):
        p = self.par(texto, 'List Number')
        if not self.lista_num_aberta:  # cada lista começa no número escrito no primeiro item
            num_id = self._nova_numeracao(inicio)
            if num_id is not None:
                pPr = p._p.get_or_add_pPr()
                numPr = OxmlElement('w:numPr')
                il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0')
                ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num_id))
                numPr.append(il); numPr.append(ni)
                pPr.append(numPr)
                self._num_atual = num_id
        else:
            if getattr(self, '_num_atual', None) is not None:
                pPr = p._p.get_or_add_pPr()
                numPr = OxmlElement('w:numPr')
                il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0')
                ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(self._num_atual))
                numPr.append(il); numPr.append(ni)
                pPr.append(numPr)
        self.lista_num_aberta = True
        return p

    def _nova_numeracao(self, inicio=1):
        try:
            estilo = self.st['List Number']
            numPr = estilo.element.pPr.find(qn('w:numPr'))
            base = numPr.find(qn('w:numId')).get(qn('w:val'))
            numbering = self.d.part.numbering_part.element
            abstract = None
            for n in numbering.findall(qn('w:num')):
                if n.get(qn('w:numId')) == base:
                    abstract = n.find(qn('w:abstractNumId')).get(qn('w:val'))
            novo = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
            num = OxmlElement('w:num'); num.set(qn('w:numId'), str(novo))
            an = OxmlElement('w:abstractNumId'); an.set(qn('w:val'), abstract); num.append(an)
            lo = OxmlElement('w:lvlOverride'); lo.set(qn('w:ilvl'), '0')
            so = OxmlElement('w:startOverride'); so.set(qn('w:val'), str(inicio)); lo.append(so); num.append(lo)
            numbering.append(num)
            return novo
        except Exception:
            return None

    def figura(self, caminho, largura, legenda):
        p = self._aplicar_quebra(self.d.add_paragraph()); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(BASE + caminho, width=Cm(largura))
        self.par(legenda, 'Caption')

    @staticmethod
    def shade(cell, fill):
        s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
        cell._tc.get_or_add_tcPr().append(s)

    def _tabela_base(self, nlin, ncol, larguras):
        tb = self.d.add_table(rows=nlin, cols=ncol); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        bd = OxmlElement('w:tblBorders')
        for side in ('top', 'bottom', 'insideH'):
            e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D3DEDB')
            bd.append(e)
        tb._tbl.tblPr.append(bd); tb.autofit = False
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
        for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
            gc.set(qn('w:w'), str(int(w * 567)))
        for r in tb.rows:  # linha não se parte entre páginas
            r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(th)
        return tb

    def tabela(self, linhas, larguras, legenda):
        ncol = max(len(r) for r in linhas)
        tb = self._tabela_base(len(linhas), ncol, larguras)
        zebra = 0
        grupos = []
        for i, row in enumerate(linhas):
            # linha de grupo: só a primeira célula, em negrito; vira faixa mesclada em Papel
            grupo = i > 0 and row and row[0].startswith('**') and row[0].endswith('**') and all(not c for c in row[1:])
            if grupo:
                grupos.append(i)
                zebra = 0
            else:
                zebra += 1
            for j in range(ncol):
                txt = row[j] if j < len(row) else ''
                c = tb.cell(i, j); c.width = Cm(larguras[j])
                p = c.paragraphs[0]; p.style = self.st['Lumis Tabela']
                if i == 0:
                    p.paragraph_format.keep_with_next = True  # cabeçalho não fica sozinho no pé da página
                    r = p.add_run(txt.replace('**', '')); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    self.shade(c, '0F2D3A')
                elif grupo:
                    self.shade(c, 'EEF7F3')
                    p.paragraph_format.keep_with_next = True  # a faixa do grupo não fica sozinha no pé da página
                    if j == 0:
                        r = p.add_run(txt.strip('*')); r.bold = True; r.font.color.rgb = RGBColor.from_string('1E8C6B')
                else:
                    self.runs(p, txt)
                    if zebra % 2 == 0:
                        self.shade(c, 'F1F6F4')
                    n = len(linhas)
                    if i == n - 1:  # a última linha leva a legenda junto (o cabeçalho já segura a 1ª linha)
                        p.paragraph_format.keep_with_next = True
        for i in grupos:  # mescla a faixa do grupo
            cel = tb.cell(i, 0).merge(tb.cell(i, ncol - 1))
            for extra in cel.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
        self.par(legenda, 'Caption')

    def tabela_niveis(self, linhas, larguras, legenda):
        """Tabela agrupada por nível (Entrega 1): célula do nível mesclada, cor por grupo e 'destrava' em itálico."""
        cab, dados = linhas[0], linhas[1:]
        tb = self._tabela_base(len(dados) + 1, 3, larguras)
        for j, txt in enumerate(cab[:3]):
            c = tb.cell(0, j); c.width = Cm(larguras[j]); p = c.paragraphs[0]; p.style = self.st['Lumis Tabela']
            p.paragraph_format.keep_with_next = True
            r = p.add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); self.shade(c, '0F2D3A')
        grupos = []
        for row in dados:
            nivel = row[0] or (grupos[-1][0] if grupos else '')
            if grupos and grupos[-1][0] == nivel and not row[0]:
                grupos[-1][1].append(row)
            elif grupos and grupos[-1][0] == nivel:
                grupos[-1][1].append(row)
            else:
                grupos.append((nivel, [row]))
        i = 1
        for nivel, itens in grupos:
            cor, fundo = NIVEIS_COR.get(nivel, ('0F2D3A', None))
            first = i
            for row in itens:
                elemento = row[1] if len(row) > 1 else ''
                porque = row[2] if len(row) > 2 else ''
                destrava = row[3] if len(row) > 3 else ''
                for j in range(3):
                    c = tb.cell(i, j); c.width = Cm(larguras[j]); c.vertical_alignment = VA.CENTER
                    c.paragraphs[0].style = self.st['Lumis Tabela']
                    if fundo:
                        self.shade(c, fundo)
                self.runs(tb.cell(i, 1).paragraphs[0], '**' + elemento + '**' if elemento and '**' not in elemento else elemento)
                self.runs(tb.cell(i, 2).paragraphs[0], porque)
                if destrava:
                    q = tb.cell(i, 2).add_paragraph(style='Lumis Tabela')
                    rr = q.add_run(destrava); rr.italic = True; rr.font.color.rgb = RGBColor.from_string('1E8C6B')
                if i == len(dados):  # a última linha leva a legenda junto
                    for cc in tb.rows[i].cells:
                        cc.paragraphs[0].paragraph_format.keep_with_next = True
                if len(itens) > 1 and i < first + len(itens) - 1:  # mantém o grupo junto
                    for cc in tb.rows[i].cells:
                        cc.paragraphs[0].paragraph_format.keep_with_next = True
                i += 1
            cel = tb.cell(first, 0)
            if i - 1 > first:
                cel = cel.merge(tb.cell(i - 1, 0))
                for extra in cel.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
            p0 = cel.paragraphs[0]; p0.style = self.st['Lumis Tabela']
            rn = p0.add_run(nivel); rn.bold = True; rn.font.color.rgb = RGBColor.from_string(cor)
            cel.vertical_alignment = VA.CENTER
        self.par(legenda, 'Caption')

    def caixa(self, linhas):
        tb = self.d.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tb.cell(0, 0); self.shade(c, 'F1F6F4')
        c.paragraphs[0].style = self.st['Lumis Tabela']; c.paragraphs[0].add_run(linhas[0] if linhas else '')
        for l in linhas[1:]:
            c.add_paragraph(l, style='Lumis Tabela')
        self.d.add_paragraph(style='Lumis Nota')

    def quebra(self):
        self.quebra_pendente = True  # vira 'quebra antes' no próximo parágrafo criado

    def _aplicar_quebra(self, p):
        if getattr(self, 'quebra_pendente', False):
            p.paragraph_format.page_break_before = True
            self.quebra_pendente = False
        return p

    def titulo(self, texto, nivel):
        return self._aplicar_quebra(self.d.add_heading(texto, level=nivel))


def montar(modo, conteudo, saida=None):
    linhas = ler_conteudo(conteudo)
    bs = blocos(linhas)
    figs, tabs, refs = numerar(bs)
    avisos = []
    D = Doc(modo) if saida else None
    for tipo, d, org in bs:
        if tipo != 'LINHA':
            pass
        if tipo == 'LINHA':
            t = d['texto']
            mnum = re.match(r'^(\d+)\.\s+(.*)$', t)
            if not mnum and D:
                D.lista_num_aberta = False
            if t.startswith('### '):
                txt = resolver(t[4:], figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.titulo(txt, 3)
            elif t.startswith('## '):
                txt = resolver(t[3:], figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.titulo(txt, 2)
            elif t.startswith('# '):
                txt = resolver(t[2:], figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.titulo(txt, 1)
            elif t.startswith('> '):
                txt = resolver(t[2:], figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt, 'Lumis Destaque')
            elif t.startswith('ROTULO:'):
                txt = resolver(t[7:].strip(), figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt, 'Lumis Rótulo')
            elif t.startswith('TITULO:'):
                D and D.par(t[7:].strip(), 'Title')
            elif t.startswith('SUBTITULO:'):
                D and D.par(t[10:].strip(), 'Subtitle')
            elif t.startswith('NOTA:'):
                txt = resolver(t[5:].strip(), figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt, 'Lumis Nota')
            elif t.startswith('LINHA:'):
                txt = resolver(t[6:].strip(), figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt, 'Lumis Tabela')
            elif t.startswith('- '):
                txt = resolver(t[2:], figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt, 'List Bullet')
            elif mnum:
                txt = resolver(mnum.group(2), figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.numerado(txt, int(mnum.group(1)))
            elif t == 'QUEBRA':
                D and D.quebra()
            elif t == 'REFERENCIAS':
                for k, n in sorted(refs.items(), key=lambda kv: kv[1]):
                    D and D.par(f'{n}. {REFS.get(k, "[referência não encontrada: " + k + "]")}', 'Lumis Nota')
            else:
                txt = resolver(t, figs, tabs, refs, avisos, org); checar_texto(txt, org, avisos)
                D and D.par(txt)
        elif tipo == 'FIG':
            D and setattr(D, 'lista_num_aberta', False)
            if d['chave'] not in FIGS:
                avisos.append(f'{org}: figura desconhecida: {d["chave"]}'); continue
            caminho, larg = FIGS[d['chave']]
            leg = f'Figura {figs[d["chave"]]}. ' + resolver(d['legenda'], figs, tabs, refs, avisos, org)
            checar_texto(leg, org, avisos)
            D and D.figura(caminho, d['largura'] or larg, leg)
        elif tipo in ('TAB', 'NIVEIS'):
            D and setattr(D, 'lista_num_aberta', False)
            if not d['linhas']:
                avisos.append(f'{org}: tabela sem linhas: {d["chave"]}'); continue
            ncol = max(len(r) for r in d['linhas'])
            larg = d['larguras']
            if tipo == 'NIVEIS':
                ncol = 3
            if len(larg) != ncol:
                avisos.append(f'{org}: tabela {d["chave"]} tem {ncol} colunas e {len(larg)} larguras')
                larg = (larg + [LARGURA_UTIL / ncol] * ncol)[:ncol]
            if sum(larg) > LARGURA_UTIL + 0.05:
                avisos.append(f'{org}: tabela {d["chave"]} com {sum(larg):.1f} cm (máximo {LARGURA_UTIL})')
            for r_i, row in enumerate(d['linhas']):
                if len(row) != len(d['linhas'][0]) and tipo == 'TAB':
                    avisos.append(f'{org}: tabela {d["chave"]}, linha {r_i + 1} com {len(row)} células (cabeçalho tem {len(d["linhas"][0])})')
            linhas_res = [[resolver(c, figs, tabs, refs, avisos, org) for c in row] for row in d['linhas']]
            for row in linhas_res:
                for c in row:
                    checar_texto(c, org, avisos, celula=True)
            leg = f'Tabela {tabs[d["chave"]]}. ' + resolver(d['legenda'], figs, tabs, refs, avisos, org)
            checar_texto(leg, org, avisos)
            if D:
                (D.tabela_niveis if tipo == 'NIVEIS' else D.tabela)(linhas_res, larg, leg)
        elif tipo == 'CAIXA':
            D and setattr(D, 'lista_num_aberta', False)
            D and D.caixa(d['linhas'])
    if saida:
        D.d.save(saida)
    return avisos, figs, tabs, refs


# ------------------------------------------------------------------ conteúdo final (v1)
CONTEUDO_DOSSIE = {
 "00_abertura.txt": "# A tese em uma página\nROTULO: Em uma frase\n> Recomendamos aceitar o aporte do Vetor Capital com liberação por etapas, ligada a condições verificáveis: primeiro tornar o ativo defensável, crescer onde a Lumis já tem cliente e contratar em ondas; novos países e plataforma, só depois. Se o Vetor exigir países ou plataforma em produção antes das condições, recomendamos não aceitar o aporte nesses termos.\nO Vetor Capital oferece R$ 120 milhões por 22% da Lumis, o que a avalia em R$ 425,5 milhões antes do aporte, cerca de 10,3 vezes a receita recorrente anual. O fundo propõe triplicar o time técnico, entrar em dois novos países e transformar o Lumis Insight em plataforma, e pede em 60 dias uma tese de crescimento defensável [fonte: Cap. 2, 1.1 e texto após o Quadro 3].\nTAB[tese]: 2.9;9.0;4.1 | Perguntas do Vetor e nossas respostas. Fonte: perguntas, Cap. 2, 1.1; respostas, leitura da equipe sobre o Anexo A.\n| Pergunta | Resposta | Onde está |\n| Onde jogar? | Hospitais e seguradoras médios no Brasil; o novo módulo de crédito, só para os bancos já clientes e só até o modo sombra (o sistema roda em paralelo, sem decidir). | Entrega 1, seção 1.4; Entrega 4, seção 4.4 |\n| Por que ali? | Clientes médios, sem IA própria, dependem mais de fornecedores externos [hipótese]; quantas das 38 contas são de porte médio [não consta]. A Aster não chega a seguradoras e bancos, 14 das 38 contas. | Entrega 1, seções 1.2 e 1.4 |\n| Contra quem? | A Aster Health, que embute a IA como módulo num sistema hospitalar usado em 210 hospitais e já instalado em boa parte dos clientes que a Lumis pretendia conquistar [fonte: Cap. 2, 1.2 e Quadro 6]; contra ela, a Lumis só tem a qualidade, e ainda precisa prová-la em campo. Depois, grandes clientes com IA própria e o fornecedor de modelo. | Entrega 1, seção 1.2 |\n| Com que ativo próprio? | Nenhum difícil de copiar hoje: a parte exclusiva da base tem contratos frágeis. A construir: dado de desfecho com direito de uso limpo e prova auditável de desempenho por grupo. | Entrega 1, seção 1.4; Entrega 2, seção 2.1 |\n| Sob que riscos? | A expansão leva quatro riscos sem solução no Brasil: equidade, privacidade, transparência e governança. O pior número: 31,8% dos idosos de CEP D/E que precisavam de prioridade não a receberam, três vezes o melhor grupo. | Entrega 5, seções 5.1 e 5.2 |\nAbrimos pela Entrega 5, que o comitê do Vetor lerá primeiro e sem a qual não aprova operação [fonte: Cap. 2, 1.1 e 4]. As outras quatro sustentam as condições dela. A Entrega 1 mostra onde a Lumis está, e a 2 mostra que o ativo ainda não passa numa auditoria independente e o que passar a medir. A 3 dá cargo a cada decisão, e a 4 escolhe o que se constrói e o que se recusa. O fechamento reúne as condições prévias, e a recomendação ao conselho segue no memorando que acompanha o dossiê.\n% A Entrega 5 começa em página nova.\nQUEBRA\n",
 "10_entrega5.txt": "# Entrega 5 · Due diligence ESG da expansão\n> A expansão leva para novos mercados quatro riscos que a Lumis ainda não resolveu no Brasil: equidade, privacidade, transparência e governança. O crescimento só se sustenta com condições prévias e com uma métrica pública, o falso negativo por grupo publicado a cada trimestre, que mostre se o dano está caindo.\nAnalisamos a expansão proposta pelo Vetor Capital (dois novos países e o Lumis Insight como plataforma) e dois casos do backlog: a expansão para o México e o módulo de risco de crédito para bancos [fonte: Cap. 2, 1.1 e Quadro 16].\n\n## 5.1 O que é material\nCada tema recebeu duas notas: o dano que o negócio pode causar a pessoas e ao ambiente, e o risco financeiro que o tema traz para a Lumis. É material o tema com nota Alta em pelo menos um dos eixos.\nFIG[e5_materialidade]: Matriz de materialidade dupla. Fonte: Cap. 2, Anexo A; notas da equipe.\nTAB[e5_temas]: 3.6;12.4 | Evidência principal de cada tema; as notas estão na Figura {fig:e5_materialidade}. Fonte: Cap. 2, Quadros 7, 9, 10, 11, 12, 14, 15 e 17.\n| Tema | Evidência principal |\n| T1 Equidade no acesso | Falso negativo de 31,8% em 60+ de CEP D/E; a faixa D/E concentra 64% das reclamações |\n| T2 Privacidade e dados de saúde | 35,4% dos dados de clientes com autorização fraca ou nenhuma; política para dados de saúde em elaboração desde 2024, sem versão aprovada |\n| T3 Transparência | 640 mil priorizações por mês com revisão humana de 2%; 94% divulgado contra 87,6% em campo |\n| T4 Governança da IA | Sem comitê de ética ou de risco; só o CTO libera versões; 2 de 4 incidentes detectados pelo cliente |\n| T5 Clima e retenção | eNPS de −31 no time técnico; rotatividade de 27% no time de dados |\n| T6 Diversidade | Mulheres são 12% e pessoas negras 6% da liderança |\n| T7 Energia | 1.240 MWh por ano, 75% em inferência; 62% de fonte renovável |\nOs temas T1 a T4 têm nota Alta nos dois eixos, já são problemas hoje e crescem com a expansão: mais clientes no mesmo modelo, populações sem histórico na base e dados de saúde cruzando fronteiras. Clima e retenção (T5) também é material, pelo risco à Lumis; a Entrega 4 trata dele. Diversidade e energia, que o capítulo inclui entre os temas de uma empresa de IA em saúde [fonte: Cap. 2, 2.7], ficam em monitoramento: nenhuma tem nota Alta, e o Anexo A não liga as duas a dano ou perda já ocorridos [não consta]. O consumo de energia, porém, cresce junto com a inferência [hipótese].\n\n## 5.2 Quem pode ser prejudicado\nO erro cresce com a idade e com a faixa de CEP (Figura {fig:e2_subgrupo}). No grupo 60+ de CEP D/E, quase um em cada três pacientes que precisavam de prioridade não a recebe, e a empresa nunca cruzou as reclamações com esse desempenho [fonte: Cap. 2, Quadros 10 e 17].\nFIG[e2_subgrupo]: Pacientes que precisavam de prioridade e não receberam, 1º semestre de 2026. Fonte: Cap. 2, Quadros 9 e 10.\nTAB[e5_grupos]: 2.5;4.7;8.8 | Grupos em risco em cada frente da expansão. Fonte: Cap. 2, Anexo A; leitura da equipe.\n| Frente | Grupo em risco | Por quê |\n| Plataforma e novos países | Pessoas de baixa renda e com pouco acesso à rede de saúde | As variáveis de acesso penalizam quem usa menos o sistema (seção 2.2) [hipótese]. Perfil de acesso no país de destino: [não consta] |\n| México | Grupos que a Lumis ainda não consegue identificar | Sem dado local de desfecho, a Lumis não sabe onde vai errar |\n| Crédito para bancos | Moradores de CEP de baixa renda e pessoas com pouco histórico financeiro | CEP e cobertura pesam 16,3% no modelo de saúde; reaproveitados, podem negar crédito por endereço [hipótese] |\nO Anexo A não traz o desempenho por sexo, raça ou cor e deficiência [não consta], e por isso não há como afirmar que esses grupos estão protegidos. Para medir o risco, usamos os indicadores 1 e 3 da Entrega 2, os sinais de alerta do protocolo e do crédito (seção 3.3) e uma condição de entrada em cada novo mercado:\nTAB[e5_indicadores]: 5.4;4.4;6.2 | Medidas de equidade, lidas a cada quinze dias, salvo a representatividade (antes da entrada; depois, anual); as reclamações, na mesa de campo da seção 4.2. Limites da equipe; linhas de base: Cap. 2, Quadros 10 e 17, e conta da equipe.\n| Medida | Limite proposto | Se passar do limite |\n| Erro por grupo: falso negativo de cada grupo e razão entre o pior e o melhor (indicador 1 da Entrega 2) | Meta: até 1,5 vez em 12 meses. Limite: 2 vezes (hoje: 3,0) | Acima de 1,5 e até 2 vezes, alerta e revisão humana de 10% dos casos. Acima de 2 vezes, restrição imediata da recomendação automática no grupo |\n| Reclamações cruzadas com o erro por grupo: parcela das reclamações dividida pela de pacientes, por faixa de CEP e idade, em 90 dias (indicador 3 da Entrega 2) | Alerta: 1,5 vez. Limite: 2 vezes (hoje, faixa D/E: 2,5, com as reclamações de 12 meses; por idade [não consta]) | A partir de 1,5 vez, revisão clínica dos casos. Acima de 2 vezes, também auditoria de equidade e investigação da causa |\n| Reversão humana por grupo no protocolo clínico (registro de revisões) | Diferença de até 2 vezes entre grupos | Revisão das variáveis de acesso |\n| Representatividade dos dados no novo mercado | Amostra suficiente para medir o erro de cada grupo | O mercado não entra em produção |\n| Aprovação no crédito por faixa de CEP | Mínimo de 0,8 da taxa da faixa com mais aprovações{ref:cfr} | Revisão humana e retirada de CEP e cobertura do modelo |\nDois grupos já passam do limite: os idosos de CEP C (2,2 vezes) e de CEP D/E (3,0 vezes). Os adultos de 18 a 59 anos de CEP D/E (1,9 vez) estão na faixa de alerta, e os idosos de CEP A/B, no limite da meta (1,5 vez) [fonte: Cap. 2, Quadro 10; conta da equipe]. Recomendamos a restrição imediata nos dois primeiros (seção 2.5). Pela linha de responsabilidade (seção 3.3), o Responsável por Dados apura as medidas da tabela, o Head of AI Management decide a revisão, a restrição ou a auditoria, e o CTO executa.\n\n## 5.3 Privacidade e dados sensíveis\nDado de saúde é dado pessoal sensível na LGPD (art. 5º, II){ref:lgpd}, como os atendimentos do Hospital Vila Ipê e da Rede Sanare [fonte: Cap. 2, Quadro 7]. Nenhum dos seis instrumentos de dados passou por revisão jurídica desde a assinatura, e o da Seguradora Prisma vence em 12/2026 (base contratual na seção 2.1); a política para dados de saúde está em elaboração desde 2024, sem versão aprovada [fonte: Cap. 2, Quadros 7 e 17].\nSeguem sem resposta no caso se a anonimização dos dados da Rede Sanare, de que depende a autorização, é irreversível (art. 12), e de onde vêm os dados sintéticos [não consta].\nTAB[e5_lgpd]: 2.4;6.4;7.2 | Pontos da LGPD que pesam na expansão. Fonte: Lei 13.709/2018; leitura da equipe, a confirmar com a DPO.\n| LGPD | O que diz | Onde a Lumis fica exposta |\n| Art. 11, § 4º | Vedado o uso compartilhado de dado de saúde entre controladores para obter vantagem econômica, salvo na prestação de serviços de saúde | Plataforma que aprende com vários clientes e modelo de crédito treinado com dado de saúde [hipótese] |\n| Art. 11, § 5º | Operadoras de planos de saúde não podem usar dado de saúde para selecionar riscos | Classificação de sinistros para 9 seguradoras, 74 mil decisões por mês [fonte: Cap. 2, Quadros 3 e 12], se alguma for operadora de saúde [hipótese] |\n| Art. 33 | Transferência internacional só nas hipóteses do artigo, como proteção adequada no destino, cláusulas contratuais ou consentimento específico | Todo novo país. Lei aplicável no México: [não consta] |\n| Art. 38 | A ANPD pode exigir relatório de impacto à proteção de dados, inclusive de dados sensíveis | A plataforma e cada novo país. Se a Lumis já tem esse relatório [não consta] |\nPara o crédito, propomos uma regra fixa: dado de saúde não entra no modelo de crédito, que deve ser treinado só com dados financeiros autorizados. Hoje, dos cinco bancos clientes, só o Banco Meridiano autoriza esse uso, com auditoria anual do cliente [fonte: Cap. 2, Quadros 3 e 7].\nTAB[e5_privacidade]: 6.0;4.0;4.0;2.0 | Privacidade: o indicador 2 da Entrega 2 e dois controles do compromisso de Segurança e Privacidade, da Declaração de Intenção da Fase 1. Responsável: DPO (seção 3.3). Base: Cap. 2, Quadros 7 e 14; conta da equipe.\n| Medida | Hoje | Meta | Frequência |\n| Direito de uso da base (indicador 2 da Entrega 2, contado em registros): registros de clientes no treino com autorização explícita e revisão jurídica | 0% com revisão jurídica; 64,6% com autorização explícita | 100% antes de qualquer frente nova; registro sem base sai do treino | Trimestral |\n| Acessos a dados de pacientes revisados em 90 dias | [não consta] | 100% | A cada 90 dias |\n| Incidentes de privacidade e dias até a ação corretiva | Sem categoria própria no registro de incidentes | Registro no dia em que é identificado; ação corretiva documentada | Mensal |\n\n## 5.4 A métrica que vamos publicar\nA Lumis não publica hoje nenhuma métrica de impacto [fonte: Cap. 2, Quadro 17]. Propomos publicar o indicador 1 da Entrega 2, sem a abertura por cliente: a taxa de falso negativo da priorização de atendimento, por grupo de idade e faixa de CEP. Foi esse o número que o Hospital Vila Ipê calculou por conta própria e enviou com a notificação [fonte: Cap. 2, Quadro 10]. Ele mede o dano mais grave do produto: o paciente que precisava de prioridade e não a recebeu.\nTAB[e5_metrica]: 3.4;12.6 | Ficha da métrica pública. Metas propostas pela equipe.\n| Item | Definição |\n| O que se publica | Falso negativo, sensibilidade, parcela marcada como prioridade e tamanho da amostra de cada grupo, mais a razão entre o pior e o melhor grupo e o valor do período anterior |\n| Ponto de partida | De 10,6% a 31,8%; razão de 3,0 vezes (1º semestre de 2026) [fonte: Cap. 2, Quadro 10]. Parcela marcada como prioridade: [não consta] |\n| Meta | Razão de até 1,5 vez em 12 meses; nenhum grupo acima de 7,4%, o falso negativo da validação declarada [fonte: Cap. 2, Quadro 9], em 24 meses |\n| Periodicidade | Trimestral, no site da Lumis e no relatório de cada cliente; leitura interna quinzenal |\n| Responsável e verificação | Head of AI Management publica; Responsável por Dados apura; auditoria independente uma vez por ano |\n| Se piorar | Acima de 2 vezes, restrição imediata da recomendação automática no grupo |\nA parcela de pacientes marcados como prioridade é publicada junto para evitar que o falso negativo caia às custas de marcar todos como prioridade: quando uma medida vira meta, ela deixa de ser boa medida (Goodhart) [fonte: Cap. 2, 2.3].\nAs condições para a expansão estão reunidas no fechamento (Tabela {tab:condicoes}).\n",
 "20_entrega1.txt": "# Entrega 1 · O mapa do território\n> O que faz o Lumis Insight funcionar é alugado, e o único ativo que poderia ser da Lumis, a base de dados, ainda não é dela. Por isso, a tese de crescimento precisa partir do que ainda vai ser construído.\n\n## 1.1 Onde a Lumis está\nA Lumis só controla a camada de cima. Nas outras três, aluga ou depende da autorização de terceiros.\nFIG[e1_camadas]: O que a Lumis controla em cada camada. Fonte: Cap. 2, Quadros 3, 4, 5 e 7; conta e leitura da equipe.\n- **Modelo:** o preço em queda alivia o custo, mas também barateia a entrada de concorrentes.\n- **Aplicação:** é da Lumis, mas quem decide o que entra no fluxo do hospital é o fornecedor do prontuário [hipótese].\n- **Integração e operação:** atravessam todas as camadas. A implementação e o treinamento junto ao cliente sustentam o valor percebido do produto [hipótese].\nPara Shapiro e Varian{ref:shapiro}, o valor fica com quem controla o gargalo. Aqui os gargalos são a distribuição dentro do hospital e o direito de usar o dado de desfecho, e nenhum dos dois está com a Lumis. Por isso a Lumis negocia em desvantagem com os fornecedores de cima, que podem mudar preço e termos, e com os clientes grandes, que podem fazer a própria IA; tem mais força com os clientes médios [hipótese].\n\n## 1.2 Quem pode tomar esse espaço\nO mercado endereçável no Brasil é estimado em R$ 2,1 bilhões em 2026 e R$ 3,4 bilhões em 2029, e a Lumis tem cerca de 2% dele [fonte: Cap. 2, texto após o Quadro 6]. A ameaça vem de quatro direções, com pesos diferentes.\nTAB[e1_ameacas]: 3.4;7.6;5.0 | As ameaças, da mais forte para a mais fraca. Fonte: Cap. 2, Anexo A; leitura da equipe.\n| Ameaça | Por que preocupa | O que joga a favor da Lumis |\n| 1. Dentro do hospital: Aster Health (direto) | Vende a IA como módulo, por R$ 340 mil a mais, sobre o contrato que o hospital já tem [fonte: Cap. 2, Quadro 6]. Para o nosso stakeholder, a pergunta passa a ser ativar o módulo do fornecedor atual ou contratar um novo. O padrão já existe no Brasil com os prontuários da MV e da Philips (Tasy){ref:mv_klas}. | A Aster não chega a seguradoras e bancos, que são 14 das 38 contas. |\n| 2. O próprio cliente: grandes clientes (potencial) | São os que mais conseguem fazer a própria IA. O Einstein, por exemplo, já usa perto de 120 algoritmos próprios{ref:einstein}. | Clientes médios, sem time próprio de IA, dependem mais de fornecedores externos [hipótese]. |\n| 3. De cima: os próprios fornecedores (potencial) | O fornecedor de modelo anunciou um módulo próprio para saúde e passa a concorrer na aplicação. Nuvem e bases clínicas podem seguir o mesmo caminho [hipótese]. | Nenhum deles tem hoje o fluxo clínico dentro do hospital [hipótese]. |\n| 4. Por fora: Núcleo Saúde Analytics (indireto) | Já tem 74 contas de BI hospitalar, quase o dobro da Lumis, e só falta a camada preditiva. | A Lumis já entrega a camada preditiva que falta à Núcleo. |\n| 4. Por fora: consultorias e integradores (indireto) | Vendem projetos sob medida aos clientes grandes, sobre a mesma infraestrutura e os mesmos modelos. | Produto pronto, com ticket médio anual de R$ 1,084 mi, contra R$ 1,5 a 4,0 mi por projeto. |\nDistribuição não garante qualidade: o modelo de sepse da Epic, embutido no prontuário, deixou de identificar dois terços dos casos{ref:wong}. A Lumis, porém, ainda não pode competir por qualidade, porque os 94% de acurácia caem para 87,6% em campo.\n\n## 1.3 O que pode mudar sem a Lumis decidir\nDuas das cinco dependências pedem atenção agora, e as duas têm prazo.\n- **Dados de clientes, a mais grave:** é a única sem substituto, e o risco já existe: o modelo atual foi treinado com os 35,4% dos dados de clientes que têm autorização fraca ou nenhuma (seções 2.1 e 5.3).\n- **Fornecedor de modelo, a mais rápida:** muda preço e termos com 30 dias de aviso. Se o preço dobrar, a margem bruta cai de 58% para 39,5%, e o caixa passa a durar 8,7 meses, em vez de 11,6 [fonte: conta da equipe sobre os Quadros 3 e 4]. Já houve queda de desempenho após uma atualização do fornecedor, em 09/2025, e há precedente de aumento: a Anthropic lançou em 2024 a nova versão do Haiku a quatro vezes o preço da anterior{ref:techcrunch}.\nTAB[e1_dependencias]: 2.7;2.9;5.7;4.7 | As cinco dependências e o que recomendamos (ainda não praticado). Fonte: Cap. 2, Anexo A; cálculos da equipe.\n| Dependência | Prazo | O que pode acontecer | Resposta recomendada |\n| Dados de clientes | Já hoje; contratos vencem em 12/2026 (Prisma) e 03/2027 (Vila Ipê) | Questionamento do uso atual obriga a retirar 35,4% dos dados de clientes e retreinar; renegociação restringe o uso; multa da LGPD de até 2% do faturamento, limitada a R$ 50 mi por infração{ref:lgpd}; clientes grandes levam os dados para a Aster. | Revisão jurídica dos seis instrumentos; aditivos de direito de uso antes dos vencimentos; linhagem de dados para isolar os registros frágeis. |\n| Modelo fundacional (44,3% do custo) | 30 dias de aviso, a qualquer momento | Margem de 54,3% com preço +20% e de 39,5% com +100%; descontinuação obriga a revalidar o sistema; fornecedor vira concorrente. | Homologar um segundo fornecedor; travar preço e versão; avaliar modelo próprio no núcleo preditivo. |\n| Câmbio (72,2% do custo em dólar) | Contínuo | Cada R$ 0,10 a mais no dólar custa R$ 19.240 por mês; a R$ 6,50, a margem vai a 51,9%. | Proteção cambial e cláusula de reajuste nos contratos com clientes. Se já há essa proteção [não consta]. |\n| Nuvem (27,9% do custo) | Renovação em 04/2027; migrar leva 7 meses | Renovação sem alternativa; cada 10% de reajuste custa R$ 40.176 por mês. | Começar já o plano de portabilidade; negociar cláusulas de saída. |\n| Bases clínicas (7,8% do custo) | Renovação anual automática | Reajuste sem teto; a renovação automática dificulta a saída. | Negociar teto de reajuste na renovação. |\n\n## 1.4 O que é difícil de copiar\nPartimos da leitura de que o diferencial está na combinação de produto, integração e conhecimento. Testada contra os dados, nenhuma parte dela é hoje difícil de copiar. As duas que mais valem, a base histórica e a validação clínica, estão travadas.\nNIVEIS[e1_copiar]: 3.6;4.2;8.2 | Cada elemento, do mais fácil ao mais difícil de copiar. Fonte: Cap. 2, Anexo A; classificação da equipe.\n| Hoje | Elemento | Por quê |\n| Copiável | Modelo e tecnologia de IA | É de terceiro, à venda por consumo, com preço em queda. |\n| Parcialmente difícil | Produto especializado | A Aster embute priorização no hospital, e os grandes clientes internalizam. |\n| Parcialmente difícil | Integração e relacionamento | Churn de 11% ao ano; a retenção não foi testada contra a Aster. |\n| Moderadamente difícil | Conhecimento e equipe | 48 técnicos e 4 anos de domínio, que saem da empresa com as pessoas. |\n| Valioso, mas travado | Base histórica | Dados de clientes com autorização fraca e viés. | Destrava com contratos regularizados e viés corrigido. |\n| Valioso, mas travado | Validação clínica | Os 94% caem para 87,6% em campo. | Destrava com validação em campo por grupo. |\n| Difícil de copiar | Nada hoje | A construir: dado de desfecho com direito de uso limpo e prova auditável de desempenho por grupo. |\nA base histórica parece o ativo mais forte, mas só os dados de clientes, 26,8% dela, são exclusivos da Lumis: o DATASUS é público, e outra empresa também consegue gerar dados sintéticos (Figura {fig:e2_base}, na seção 2.1). Mesmo esses dados têm autorização fraca ou nenhuma em 35,4% dos casos e, usados como estão, tendem a repetir o viés da seção 5.2 [hipótese]. Há mais duas travas:\n- **Concentração:** só 2 dos 24 clientes de saúde fornecem dados, e cerca de três quartos do dado clínico de clientes vêm de um único cliente, a Rede Sanare.\n- **Escala:** quem fornece o prontuário, como a Aster com 210 hospitais, pode ter acesso a muito mais dados [hipótese].\nO que pode valer mais que a experiência da equipe nos 4 anos de operação é o histórico dos dados de clientes, de 2019 a 2026, porque pode registrar a priorização clínica brasileira junto com o desfecho, algo que nenhum fornecedor de modelo tem [hipótese]. Esse histórico só vira barreira com direito de uso limpo e prova auditável de desempenho por grupo. Para chegar lá:\n- Regularizar os contratos de dados antes dos vencimentos (12/2026 e 03/2027) e estender o direito de uso aos demais clientes.\n- Trocar os 94% por validação em campo, por grupo e publicada (seção 5.4).\n- Focar em hospitais e seguradoras médios, onde a integração vale mais e o cliente não faz a própria IA [hipótese]. O novo módulo de crédito para os bancos já clientes é exceção ao foco, só até o modo sombra (seção 4.4).\n",
 "30_entrega2.txt": "# Entrega 2 · Auditoria do ativo: dados e métricas\n> Hoje o ativo da Lumis não passa numa auditoria independente. A parte da base que só a Lumis tem está sob contratos frágeis, e o modelo erra mais justamente com quem mais precisa de prioridade.\n\n## 2.1 De onde vêm os dados e se a Lumis pode usá-los\nFIG[e2_base]: A base de treinamento do Lumis Insight. Fonte: Cap. 2, Quadro 7.\n- **Só pouco mais de um quarto da base é exclusivo da Lumis.** O DATASUS, a maior parte, está aberto a qualquer concorrente. Dos dados de clientes, 35,4% (2,09 milhões de registros) vêm de contratos que não autorizam o treino com clareza, e esses são os primeiros a vencer.\n- **O risco já existe.** Nenhum dos seis instrumentos que dão base aos dados, quatro deles contratos com clientes, passou por revisão jurídica desde a assinatura, e o modelo atual já foi treinado com esses dados. Dado de saúde é dado sensível, e a LGPD pede finalidade específica e informada (arts. 5º, 6º e 11){ref:lgpd}. Quem trata dados em nome do cliente só pode usá-los para a finalidade definida por ele{ref:anpd}, e uma cláusula de \"melhoria contínua do serviço\" dificilmente cobre treinar um produto vendido a outros clientes.\n- **Parte dos dados é anterior à empresa.** Há registros desde 2019, e a Lumis foi fundada em 2022 [fonte: Cap. 2, Quadros 3 e 7]. Sob que cláusula esse histórico chegou [não consta].\n**Recomendação:** parecer da DPO sobre os seis instrumentos antes do próximo treino, e renegociação de Prisma e Vila Ipê antes do vencimento (12/2026 e 03/2027), incluindo o histórico.\n\n## 2.2 O que o modelo mede de fato\nAs duas variáveis que mais pesam no modelo deveriam medir necessidade e gravidade. Na prática, medem acesso: quem conseguiu ser atendido e quanto isso custou. Quem teve menos acesso parece menos grave. É o problema da variável substituta (proxy), e \"foi exatamente isso que aconteceu no incidente da Fase 1\" [fonte: Cap. 2, 2.2].\nFIG[e2_proxy]: Quatro variáveis (49,8% do peso do modelo) medem acesso. Fonte: Cap. 2, Quadro 8; conta e leitura da equipe.\nDas outras seis variáveis, a idade (12,7%) mede o que diz medir. Comorbidades (11,2%) e exames (6,9%) só captam o que foi diagnosticado ou pedido, e tempo até o exame (9,8%), faltas (5,3%) e especialidade de origem (4,3%) também dependem da fila e do acesso à rede [hipótese]. Por que o erro é maior nos idosos [não consta].\n**Um caso real com o mesmo erro.** Em 2019, um estudo publicado na revista Science mostrou que um algoritmo usado nos Estados Unidos previa o custo do paciente, e não a doença{ref:obermeyer}. Pacientes negros, com menos acesso, gastavam menos e recebiam o mesmo escore de pacientes brancos mais saudáveis.\n\n## 2.3 O que acontece em campo\nOs 94% de acurácia vêm de uma amostra de 2023 com dois hospitais da mesma região. Em campo, a parcela de pacientes que precisavam de prioridade e não receberam mais que dobrou, e o erro se concentra em quem já é mais vulnerável (Figura {fig:e2_subgrupo}, seção 5.2) [fonte: Cap. 2, Quadros 9 e 10].\nQuem mediu o pior número foi o Hospital Vila Ipê, e não a Lumis [fonte: Cap. 2, Quadro 10].\n\n## 2.4 Os números que mostramos ao mercado\nO Vetor Capital avisou: número que não sobrevive a uma auditoria independente vira passivo. Passamos as métricas de hoje por quatro perguntas (Figura {fig:e2_metricas}) [fonte: Cap. 2, 2.3], e nenhuma fica como está. O caso mais grave são os \"5 milhões de vidas\": a contagem repetida foi achada em janeiro e continua no material comercial [fonte: Cap. 2, Quadros 11 e 14]. Manter esses números contradiz o compromisso de Transparência, da Declaração de Intenção da Fase 1.\nFIG[e2_metricas]: As cinco afirmações do painel comercial e um indicador que descartamos (auditorias de contrato), com as quatro perguntas. Todos saem ou são trocados. Fonte: Cap. 2, Quadros 9, 11 e 14; leitura da equipe.\nO que cada número permite concluir: os 94% valem para a validação de 2023 (em campo, 87,6%); os 30% vêm de um piloto de seis semanas sem grupo de controle, que não mostra se a redução veio do produto; os 5,2 milhões contam registros, com o mesmo paciente repetido; o NPS 72 vale para 9 respondentes escolhidos pelo comercial; e os 99,92% cobrem só a interface de programação [fonte: Cap. 2, Quadros 9 e 11].\n\n## 2.5 O que passamos a medir e o que decidimos agora\nNo lugar dessas métricas, propomos cinco indicadores (Figura {fig:e2_indicadores}). Cada um muda uma decisão concreta e, quando cabe, se abre por grupo. Como a Lumis mede o erro do próprio produto, cada um tem um conferente que não produziu o número. Quem apura segue a linha da seção 3.3: o Responsável por Dados apura os indicadores 1 e 3, a DPO o 2 e o Head of AI Management o 4 e o 5.\n**A decisão de agora.** Dois grupos já passam do limite do indicador 1 (2 vezes o erro do melhor grupo): idosos de CEP C e de CEP D/E. Diante de padrão de viés, o compromisso de Não Amplificação de Danos prevê investigar a causa, corrigir e, quando necessário, suspender o uso até a correção. Por isso recomendamos restringir já a recomendação automática nesses grupos: o Head of AI Management ordena e o CTO executa (seção 3.3). Cerca de um quarto das priorizações, por volta de 154 mil por mês [hipótese: distribuição das decisões igual à da base], volta à triagem do hospital, e o modelo segue rodando em paralelo, sem decidir, para medir a correção. A carga em cada hospital e a receita afetada [não consta].\nFIG[e2_indicadores]: Os cinco indicadores, com limites propostos pela equipe. Fonte: Cap. 2, Quadros 7, 10, 11, 14 e 17.\n",
 "40_entrega3.txt": "# Entrega 3 · A linha de responsabilidade\n> Os registros da Lumis não mostram, para nenhuma decisão do Lumis Insight, um cargo que autorizou o uso, um que acompanha o erro e um que pode desligar o sistema em prazo conhecido. A única porta com dono é a liberação de novas versões, e os incidentes dos últimos dois anos entraram por outras portas [hipótese].\nO sistema toma cerca de um milhão de decisões por mês [fonte: conta da equipe sobre o Quadro 12]. Propomos um cargo para cada etapa: autorizar, acompanhar, suspender e religar.\n\n## 3.1 As decisões que o sistema toma hoje\nClassificamos como de alto impacto toda decisão que afeta a saúde ou o acesso a um serviço essencial. Volume alto e revisão já existente não rebaixam o nível; só mudam o controle necessário. Decisão sem erro medido por grupo fica como alta até ser medida. O roteamento sobe de nível se o canal receber contestação de decisão do sistema.\nTAB[e3_decisoes]: 4.8;1.6;3.5;4.5;1.6 | Decisões delegadas ao sistema. Fonte: Cap. 2, Quadro 12. Depois do plano de contenção, há revisão obrigatória nos casos de maior impacto [fonte: Cap. 2, 1.1]; quais são esses casos [não consta], assim como a aprovação interna no protocolo, no sinistro e no crédito. A coluna Impacto é classificação da equipe.\n| Decisão | Por mês | Revisão humana hoje | Quem autorizou o uso | Impacto |\n| Priorização da fila de atendimento | 640 mil | Amostra de 2% | Diretoria comercial do cliente, sem aprovação interna formal | Alto |\n| Sugestão de protocolo clínico | 210 mil | Sempre, pelo médico | Comitê clínico do hospital | Alto |\n| Classificação de risco de sinistro | 74 mil | Só acima de R$ 50 mil | Diretoria da seguradora | Alto |\n| Sinalização de risco de crédito | 31 mil | Sempre | Comitê de crédito do banco | Alto |\n| Roteamento de mensagens de suporte | 95 mil | Nenhuma | Operação da Lumis | Baixo |\nO controle segue o costume de cada cliente: o crédito, com 31 mil decisões, tem revisão em todos os casos; a priorização, com 640 mil e viés já medido, só em 2%.\n\n## 3.2 Quem responde hoje e o que os incidentes mostram\nA Lumis não tem comitê de ética nem de risco. Dos sete papéis com dono, só um controla o que entra em produção: a decisão de liberar uma nova versão do modelo é exclusiva do CTO [fonte: Cap. 2, Quadros 12 e 13]. Nenhum registro diz quem pode suspender o sistema nem em quanto tempo [não consta]. Na Fase 1, esse poder ficou com a liderança em conjunto, sem prazo [fonte: Mapa de Stakeholders, Fase 1].\nTAB[e3_incidentes]: 1.7;5.6;2.9;2.5;3.3 | Incidentes dos últimos 24 meses. Fonte: Cap. 2, Quadro 14. A coluna Por onde entrou é leitura da equipe.\n| Quando | O que aconteceu | Quem percebeu | Correção | Por onde entrou |\n| 03/2025 | Exames de um laboratório novo passaram a ser descartados | Cliente | 22 dias | Dados de entrada |\n| 09/2025 | Queda de desempenho após atualização do fornecedor do modelo | Equipe interna | 6 dias | Fornecedor |\n| 01/2026 | Registros duplicados inflaram as \"vidas analisadas\" | Auditoria interna | Só no sistema | Material comercial |\n| 07/2026 | Viés etário e regional na priorização | Hospital Vila Ipê | Em tratamento | Acompanhamento por grupo |\nSe a atualização do fornecedor de 09/2025 passou pela liberação do CTO [não consta].\n\n## 3.3 A linha de responsabilidade proposta\nTrês regras orientam o desenho. Quem libera uma versão não é o único que acompanha o erro nem o único que pode suspender. Suspender exige um só cargo, e religar exige dois. A Lumis passa a autorizar cada uso, além do cliente.\nTAB[e3_linha]: 5.6;4.6;5.8 | Quem responde por cada etapa. Proposta da equipe sobre os cargos do Cap. 2, Quadro 13.\n| Etapa | Cargo responsável | Contrapeso |\n| Autorizar cada uso em cada cliente | CEO | Parecer obrigatório do Head of AI Management e da DPO |\n| Liberar versão nova ou atualização do fornecedor | CTO | Head of AI Management confere o teste por grupo e pode barrar |\n| Medir o erro por grupo e cruzar as reclamações | Responsável por Dados | Auditor externo, contratado pela CEO, refaz a conta |\n| Acompanhar e confirmar o alerta | Head of AI Management | DPO como suplente |\n| Suspender ou restringir | Head of AI Management | CTO e DPO também podem suspender |\n| Executar a suspensão | CTO | Responsável por Dados como suplente |\n| Religar | CTO e Head of AI Management, juntos | Sem acordo, segue suspenso e a CEO leva o caso ao conselho |\n| Aprovar os números divulgados ao mercado | Head of AI Management | Responsável Comercial corrige o material |\n| Proteger dados pessoais e direito de uso | DPO | Responsável por Dados retira o dado sem autorização |\n| Prestar contas ao conselho | CEO | Head of AI Management relata as suspensões diretamente |\nA linha vale para as quatro decisões de alto impacto. Entre elas, muda só o sinal que dispara o alerta, lido a cada quinze dias, como prevê o compromisso de Não Amplificação de Danos, da Declaração de Intenção da Fase 1. Qualquer queixa de cliente, contestação de paciente ou incidente dispara o alerta na hora.\nTAB[e3_alertas]: 3.6;12.4 | Alertas por decisão. Limites nas seções 2.5 e 5.2; o do sinistro é proposta da equipe. Dos quatro sinais, o Anexo A só traz o erro por grupo da priorização, medido pelo hospital [fonte: Cap. 2, Quadro 10].\n| Decisão | Sinal de alerta |\n| Priorização da fila | Grupo acima de 1,5 vez o erro do melhor (revisão de 10%); acima de 2 vezes, restrição |\n| Protocolo clínico | Grupo com mais de 2 vezes as sugestões alteradas pelo médico |\n| Risco de crédito | Faixa de CEP com aprovação abaixo de 0,8 da melhor |\n| Risco de sinistro | Faixa de CEP com mais de 2 vezes as negativas da melhor |\n**Prazo para suspender.** Confirmado o alerta, a decisão sai em até 24 horas e a execução em mais 24, sem esperar nova leitura. Com dano clínico em curso, a execução é imediata. O teto de 48 horas é o mesmo prazo que já prometemos para revisar um caso contestado. Como o tempo real de suspensão [não consta], o prazo é uma meta até o primeiro teste. Se o teste passar de 48 horas, o teto continua e o CTO corrige o procedimento até cumpri-lo.\n**O que fazemos já:**\n- Executar já a restrição da seção 2.5: o Head of AI Management ordena e o CTO executa.\n- Subir para 10% a revisão humana no grupo de 18 a 59 anos de CEP D/E, que está na faixa de alerta (seção 5.2).\n- Testar a suspensão em até 30 dias, por cliente e por grupo, e repetir a cada trimestre. O Head of AI Management conduz o teste e o CTO executa.\n- Exigir, nos contratos, um cargo do lado do cliente que assine a autorização de uso. A cláusula fica com a DPO.\n\n## 3.4 Quando a revisão humana é obrigatória\n- **Priorização em grupo com mais de 2 vezes o erro do melhor:** todos os casos voltam à triagem do hospital. Critério: o compromisso de Não Amplificação de Danos (seção 2.5). Acima de 1,5 e até 2 vezes, a revisão sobe de 2% para 10% (indicador 1; regra na seção 5.2), para que a leitura quinzenal tenha casos suficientes.\n- **Toda contestação de paciente, profissional ou cliente:** imediata na urgência e em até 48 horas nos demais casos. Critério: o compromisso de Reversibilidade, da Declaração de Intenção da Fase 1.\n- **Toda sugestão de protocolo clínico:** como já ocorre, agora com registro de aceite, alteração ou rejeição pelo médico. Critério: o dano é à saúde, e revisão sem registro não pode ser auditada.\n- **Classificação de sinistro que leve a negativa ou atraso:** no lugar do corte de R$ 50 mil. Critério: o efeito sobre a pessoa.\n- **Toda sinalização de risco de crédito:** como já ocorre. Critério: acesso a serviço essencial.\n- **Fonte de dados nova, cliente novo ou versão nova:** revisão de 10% dos casos até uma leitura quinzenal sem piora em nenhum grupo. Critério: os incidentes de 03/2025 e 09/2025 começaram assim.\n\n## 3.5 O que muda num setor novo\nNum setor novo a Lumis não sabe reconhecer o erro típico nem tem dado local para medi-lo. O backlog já traz um pedido assim, o módulo veterinário (recusado na seção 4.5), e o Vetor Capital propõe transformar o produto em plataforma e abrir dois países [fonte: Cap. 2, 1.1 e Quadro 16]. A estrutura muda em cinco pontos:\n- **Portão de entrada.** O Responsável por Produto leva o pedido, a CEO decide e o Head of AI Management e a DPO dão parecer obrigatório. Os critérios ficam no funil da seção 4.3.\n- **Modo sombra.** O sistema roda em paralelo, sem decidir, até haver erro medido por grupo no novo domínio; a saída segue os critérios do funil (seção 4.3).\n- **Responsável do domínio.** Um especialista do setor, do cliente ou contratado, assina o que conta como erro grave, por exemplo um médico-veterinário.\n- **Suspensão testada antes de começar.** O Head of AI Management conduz o teste por mercado e por cliente, e o CTO executa.\n- **Quem mede o erro em cada mercado.** O Responsável por Dados indica a pessoa antes da entrada, e o Head of AI Management aprova.\nSem uma linha de responsabilidade explícita, a expansão para novos setores \"multiplica exposição sem multiplicar controle\" [fonte: Cap. 2, 2.4].\n",
 "50_entrega4.txt": "# Entrega 4 · Cultura e funil de inovação\n> O técnico e o comercial discordam sobre quanto tempo validar, mas partem da mesma crença: a qualidade se prova antes de o produto ir a campo, pela média. Propomos um único critério para decidir o que se promete e o que se constrói: o desempenho medido em campo por grupo de pacientes. Com ele, recusamos três dos sete pedidos do backlog.\nO backlog soma 97 meses-pessoa, mais de três vezes os 30 disponíveis no semestre, e não tem critério de risco, de prova ou de encerramento [fonte: Cap. 2, Quadro 16].\n## 4.1 A cultura em três camadas\nUsamos as três camadas de Schein{ref:schein}: o que se vê, o que se diz e o que o grupo dá como óbvio sem perceber. O técnico vê a promessa longe da entrega e se diz pouco ouvido; o comercial vê o contrário. A rotatividade em doze meses foi de 19% na empresa e de 27% no time de dados [fonte: Cap. 2, Quadro 15].\nFIG[e4_clima]: Concordância por grupo; pesquisa de julho de 2026 com 81 de 96 colaboradores. Fonte: Cap. 2, Quadro 15.\nTAB[e4_camadas]: 2.8;13.2 | A cultura nas três camadas. Fonte: Cap. 2, seções 1.2 e 2.5 e Quadros 9, 10, 11, 14, 16 e 17; leitura da equipe.\n| Camada | O que encontramos |\n| Artefatos | O número de vitrine (94%) segue no material, contra 87,6% em campo; o erro de 01/2026 seguiu no material comercial; quem mediu o erro por grupo foi o hospital; o backlog registra quem pediu, e não o risco. |\n| Valores declarados | O técnico diz que o comercial promete o que o produto não entrega. O comercial diz que o técnico atrasa tudo em nome de um rigor que o cliente não pede. |\n| Pressupostos | Cada área acredita vender uma coisa: confiabilidade ou velocidade [fonte: Cap. 2, 2.5]. Por baixo das duas, a mesma crença: \"aqui, modelo validado é modelo bom\". |\n**O pressuposto profundo** [hipótese]: \"Aqui, modelo validado é modelo bom: a qualidade se prova antes de o produto ir a campo, pela média, e depois disso o número vale.\" Se a qualidade se resolve antes, só resta discutir quanto tempo gastar antes, e sem número de campo por grupo a discussão não tem árbitro. Por isso ela nunca chegou \"para uma sala onde ela poderia ser resolvida\" [fonte: Cap. 2, 1.2].\nA crença comum explica o que a leitura de cada área não explica: as duas concordam em não saber a quem escalar, e nenhuma mediu o erro por grupo antes do hospital. Explica também por que o comercial tem mais voz: 63% do comercial dizem que a área é ouvida nas decisões de produto, contra 34% do técnico. Depois do lançamento, a única informação nova que circula é a do cliente que compra, e quem a traz é o comercial [hipótese]. Inferimos o pressuposto de quadros, sem entrevistas; a mesa de campo e a pesquisa de 01/2027 servem para testá-lo.\n## 4.2 Duas intervenções\nComunicado interno quase nunca resolve pressuposto [fonte: Cap. 2, 2.5]. As duas intervenções mudam o que se mede e quem decide.\nTAB[e4_intervencoes]: 7.6;2.9;2.6;2.9 | As duas intervenções; responsáveis, prazos e metas da equipe. Base: Cap. 2, Quadros 11 e 17; seção 2.4.\n| Intervenção | Responsável | Prazo | Sinal de que funcionou |\n| **A. Regra única de prova.** Nenhum número de desempenho sai da empresa sem o resultado de campo aberto por grupo, com amostra e data. Vale para material, proposta, aditivo e liberação de versão. Incidente só fecha com o material corrigido. | Head of AI Management; o Responsável Comercial aplica | 30 dias para as cinco afirmações do painel | Afirmações públicas auditáveis: de 0 para 5 de 5, ou retiradas |\n| **B. Mesa de campo quinzenal.** Técnico, comercial, dados, produto, DPO e Head of AI Management leem juntos o erro por grupo e as reclamações por faixa de CEP e idade (indicadores 1 e 3 da Entrega 2). A mesa propõe o prazo de validação, que o CTO decide, e responde com data aos pedidos do comercial. | Responsável por Produto | Primeira reunião em 15 dias; prazos de validação publicados em 60 | Reclamações cruzadas com o desempenho: de nunca para toda quinzena |\nA regra produz o número por grupo, e a mesa é onde ele vira decisão de produto e de venda. Ela usa o horário da leitura quinzenal que o compromisso de Não Amplificação de Danos já prevê (seção 3.3) e é a semente do comitê de risco e ética que a Lumis não tem [fonte: Cap. 2, texto após o Quadro 12], com parecer do Head of AI Management na parte de risco. As duas áreas ganham algo: o técnico, prazo de validação com regra; o comercial, prazo conhecido e um número que resiste a uma due diligence. Para medir a percepção, repetimos as afirmações da Figura {fig:e4_clima} numa pesquisa curta em 01/2027, aplicada por terceiro. Esperamos pelo menos 10 pontos a mais no técnico e em \"sei a quem escalar\" (hoje cerca de 26% de quem respondeu) [fonte: conta da equipe sobre o Quadro 15]. O resultado é julgado na pesquisa de 07/2027.\n## 4.3 Funil de inovação\nCada passagem é um portão com quatro saídas: seguir, encerrar, pausar com data ou voltar, como no Stage-Gate{ref:cooper}. Os critérios são eliminatórios, sem nota somada, para que receita não compense risco aos pacientes. O portão segue a seção 3.5 nas iniciativas que mudam decisão sobre pessoas, mercado ou setor; pipeline, explicabilidade e ISO/IEC 42001 seguem a priorização do Responsável por Produto, com parecer do Head of AI Management.\nTAB[e4_funil]: 2.2;3.8;7.1;2.9 | Critérios de cada etapa. Fatias e número de leituras: proposta da equipe; limites de erro das Entregas 2, 3 e 5.\n| Etapa | Entra quando | Avança quando | Encerra quando |\n| Entrada | Ficha com problema, cliente nomeado, base legal dos dados e métrica de sucesso | Cliente no foco da Entrega 1 ou exceção registrada; nenhuma dependência nova de fornecedor único; nada contra os compromissos; responsável com cargo | Falha em qualquer critério |\n| Descoberta (até 10% do esforço) | Aprovada na Entrada | Cliente com intenção de compra e dados autorizados por escrito; amostra para medir o erro em todo grupo; responsável do domínio nomeado | A fatia ou o prazo acabam e falta algum item |\n| Modo sombra (até mais 30%) | Roda em paralelo, sem decidir | Quatro leituras quinzenais com o pior grupo até 1,5 vez o erro do melhor (no crédito, aprovação por faixa de CEP de pelo menos 0,8 da melhor) | Seis leituras com algum grupo acima de 2 vezes |\n| Produção limitada | Condições de frente nova (Tabela {tab:condicoes}); autorização da CEO por cliente | Leitura sem piora em nenhum grupo; número publicado | Sem correção em seis leituras |\n| Escala | Auditoria independente contratada | Indicadores contínuos | Suspensão pela seção 3.3 |\nEm cada etapa, o comercial só diz o que ela já provou: \"em estudo\" na Entrada e na Descoberta, \"em teste\", sem número, no modo sombra, e o erro por grupo medido em campo na produção limitada e na escala.\n- **Quem pediu não decide.** Os pedidos do sócio-fundador, do investidor e do comercial passam pelo mesmo portão. A CEO pode decidir contra um parecer contrário, mas só por escrito, e o caso vai ao conselho.\n- **Recurso por etapa e encerramento definido antes.** Nenhuma iniciativa passa de 10 meses-pessoa sem nova decisão. Encerra a que gastar a fatia sem evidência nova, não obtiver o direito de uso dos dados até a data marcada ou tiver a estimativa de esforço aumentada em mais de 50%.\n## 4.4 Distribuição dos 30 meses-pessoa\nProporção fixa não funciona aqui: na divisão 70/20/10 de Nagji e Tuff{ref:nagji}, a faixa adjacente teria 6 meses-pessoa, sem nenhuma iniciativa adjacente que caiba inteira (a menor pede 11). Financiamos etapas, e a proporção resulta da ordem de prioridade: compromisso já assumido, depois o ativo que a seção 1.4 aponta, depois a receita por esforço.\nTAB[e4_distribuicao]: 3.4;10.2;2.4 | Distribuição por horizonte. Proposta da equipe sobre o Cap. 2, Quadro 16.\n| Horizonte | O que recebe | Meses-pessoa |\n| Melhoria do produto atual | Compromissos já assumidos 3,0; correção do viés 3,5; explicabilidade 9,0; primeira fatia do pipeline 7,2; diagnóstico da ISO/IEC 42001 0,7 | 23,4 (78%) |\n| Expansão adjacente | Crédito para bancos, até o modo sombra | 5,6 (19%) |\n| Transformação | Descoberta da prova de desempenho por grupo como oferta | 1,0 (3%) |\n- **Melhoria.** O erro de 31,8% e os compromissos das Entregas 2, 3 e 5 não têm linha no backlog, e expandir antes de resolver a base de dados é \"replicação de um problema em escala maior\" [fonte: Cap. 2, 2.6]. Não sabemos se a capacidade já desconta esse trabalho [não consta] e tratamos como se não descontasse. A primeira fatia do pipeline traz o rastreio da origem dos dados e o erro por grupo automático, e a explicabilidade dá conteúdo à revisão humana da seção 3.4.\n- **Adjacência.** O crédito tem a maior receita por esforço, R$ 0,6 mi por mês-pessoa [fonte: conta da equipe sobre o Quadro 16], 5 bancos já clientes e um terreno onde a Aster não chega. É exceção registrada ao foco da Entrega 1 e só vai à produção com as condições de frente nova (Tabela {tab:condicoes}).\n- **Transformação.** Um mês-pessoa para testar se a prova de desempenho por grupo, o ativo da seção 1.4, vira oferta. É a nossa resposta à plataforma que o Vetor propõe [fonte: Cap. 2, 1.1].\n## 4.5 O que recusamos\nTAB[e4_recusas]: 3.0;2.5;2.6;7.9 | As três recusas; as demais iniciativas recebem recurso na Tabela {tab:e4_distribuicao}. Fonte: Cap. 2, Quadro 16; decisões da equipe.\n| Iniciativa | Quem pediu | Decisão | Por quê |\n| Módulo veterinário | Sócio-fundador | Recusado | Domínio cujos dados e erros a Lumis não conhece, sem responsável do domínio, fora do foco. Ocuparia 53% da capacidade por 11% da receita listada |\n| Agente de triagem | Comercial | Recusado neste ciclo | Automatizaria de novo a etapa do erro de 31,8%. Volta ao portão com o pior grupo até 1,5 vez o melhor e a explicabilidade pronta |\n| México | Investidor | Fora desta janela | Ocuparia 73% da capacidade, sem dado local para medir o erro por grupo. Reavaliado ao fim da janela de seis meses |\nAs três recusas vêm de lados diferentes (sócio-fundador, comercial e investidor) e abrem mão de R$ 11,3 mi dos R$ 19,7 mi de receita potencial em 12 meses do backlog (57,4%) [fonte: conta da equipe sobre o Quadro 16], um teto sem garantia de realização. O veterinário só volta com parceiro de dados autorizados, responsável do domínio e o erro por grupo dentro do limite no Brasil.\nTriplicar o time técnico, como propõe o Vetor, antes de as intervenções rodarem leva gente nova para o mesmo pressuposto [hipótese]; por isso propomos contratar em ondas. Há ainda um risco de pessoa-chave: o Responsável por Dados aparece em quase todos os portões, num time com rotatividade de 27% [fonte: Cap. 2, Quadro 15].\n",
 "60_fechamento.txt": "# O que precisa estar resolvido antes\nDas cinco entregas saem doze condições prévias: cinco antes de fechar com o Vetor, uma no vencimento dos contratos de dados e seis antes de cada frente nova.\nTAB[condicoes]: 8.4;4.8;2.8 | Condições prévias; entre parênteses, a seção de onde vêm. Cargos e prazos propostos pela equipe.\n| Condição | Quem responde | Prazo |\n| **Em até 30 dias, antes de fechar com o Vetor** | | |\n| 1. Recomendação automática restrita para idosos de CEP C e D/E, com o modelo em paralelo (2.5 e 3.3) | Head of AI Management ordena; CTO executa | Já |\n| 2. As cinco afirmações do painel comercial retiradas ou tornadas auditáveis pela regra única de prova (2.4 e 4.2) | Head of AI Management; Responsável Comercial aplica | 30 dias |\n| 3. Linha de responsabilidade aprovada e suspensão testada por cliente e por grupo (3.3) | CEO aprova; Head of AI Management conduz o teste; CTO executa | 30 dias; depois, a cada trimestre |\n| 4. Mesa de campo quinzenal funcionando, com o erro por grupo cruzado com as reclamações (4.2 e 5.2) | Responsável por Produto; o cruzamento é do Responsável por Dados | Primeira reunião em 15 dias |\n| 5. Parecer da DPO sobre os seis instrumentos de dados; registro sem base sai do treino (2.1 e 5.3) | DPO | Antes do próximo treino |\n| **Antes do vencimento dos contratos** | | |\n| 6. Aditivos de direito de uso da Prisma e do Vila Ipê, incluindo o histórico (1.4 e 2.1) | DPO | Prisma: 12/2026. Vila Ipê: 03/2027 |\n| **Antes de cada frente nova** | | |\n| 7. Política de privacidade para dados de saúde aprovada; previsão de aprovação [não consta] (5.3) | DPO | Antes de qualquer frente nova |\n| 8. Erro por grupo dentro do limite no Brasil e métrica pública de falso negativo por grupo já publicada (5.2 e 5.4) | Head of AI Management; Responsável por Dados apura | Antes da produção |\n| 9. Validação local por grupo, com amostra suficiente, e lei de dados do país conferida (3.5, 5.2 e 5.3) | Responsável por Dados indica quem mede; Head of AI Management aprova; DPO confere a lei | Antes da entrada; depois, anual |\n| 10. Relatório de impacto à proteção de dados (LGPD, art. 38{ref:lgpd}) da plataforma e de cada país (5.3) | DPO | Antes da produção |\n| 11. No novo módulo de crédito, dado de saúde fora do modelo e aprovação por faixa de CEP de pelo menos 0,8 da melhor (4.3, 5.2 e 5.3) | DPO; Responsável por Dados mede a aprovação; Head of AI Management decide | Desde o primeiro dia |\n| 12. Portão de entrada; modo sombra até quatro leituras quinzenais com o pior grupo até 1,5 vez o melhor; responsável do domínio e suspensão testada no mercado (3.5 e 4.3) | CEO decide, com parecer do Head of AI Management e da DPO; teste como na condição 3 | Antes da entrada |\nAs mesmas condições marcam as etapas de liberação do aporte e de cada uso dele; o valor de cada etapa e o custo das contratações em ondas [não consta]. Se o Vetor exigir países ou plataforma em produção antes delas, recomendamos não aceitar o aporte nesses termos.\nTAB[aporte]: 3.6;6.4;6.0 | Uso do aporte. Fonte: Cap. 2, 1.1 (proposta do Vetor); recomendação da equipe.\n| O que o Vetor propõe | O que recomendamos | Condição para avançar |\n| Triplicar o time técnico | Contratar em ondas | Condições 2 e 4; ondas revistas em 01/2027, com a pesquisa curta de clima |\n| Abrir operação em dois novos países | Fora desta janela de seis meses; o México é reavaliado ao fim dela, e o segundo país [não consta] | Condições 1 a 12 |\n| Transformar o Lumis Insight em plataforma | Só 1,0 mês-pessoa para testar se a prova de desempenho por grupo vira oferta (Descoberta, seção 4.3) | Portão do funil (seção 4.3); produção só nas condições dos países |\nOs cinco compromissos da Declaração de Intenção da Fase 1 continuam com o texto atual e ganham cargo ou limite (entre parênteses, a seção):\n- **Transparência:** ganha cargo, o Head of AI Management, com a regra única de prova (4.2), a aprovação dos números divulgados ao mercado (3.3) e a métrica pública trimestral (5.4).\n- **Não Amplificação de Danos:** ganha limite e cargo, com restrição imediata acima de 2 vezes, sem esperar nova leitura (2.5 e 5.2), leitura quinzenal com alerta pelo Responsável por Dados e pelo Head of AI Management (3.3) e modo sombra antes de frente nova (3.5 e 4.3).\n- **Reversibilidade:** contestação revista em até 48 horas e registro de aceite e alteração no protocolo (3.4); as 48 horas viram também o teto da suspensão (3.3).\n- **Responsabilidade e Prestação de Contas:** o poder de suspender passa ao Head of AI Management, que já coordenava a investigação de incidentes; na Fase 1, era da liderança em conjunto, sem prazo (3.2). A CEO autoriza cada uso (3.3) e decide no portão do funil o que entra (4.3).\n- **Segurança e Privacidade:** ganha cargo, a DPO, com acessos revisados a cada 90 dias, incidentes de privacidade medidos e dado de saúde fora do crédito (5.3).\nA restrição imediata para idosos de CEP C e D/E (seção 2.5) e as três recusas do backlog (seção 4.5) seguem a regra que assumimos na Fase 1: no conflito entre a continuidade do negócio e a segurança dos afetados, a segurança vem primeiro [fonte: Mapa de Stakeholders, Fase 1].\n",
 "80_referencias.txt": "# Referências\nREFERENCIAS\nNOTA: Dados da Lumis: Cap. 2, Anexo A. Cálculos, classificações, limites, metas, prazos e cargos propostos são da equipe. As empresas, os estudos e as normas reais citados são contexto externo, usados como análogo ou referência, e não fazem parte do caso Lumis.\nNOTA: A IA generativa (Claude Code, Anthropic) apoiou a pesquisa, a conferência das fontes e a redação, com revisão da equipe. O uso na pesquisa de mercado da Entrega 1 está documentado no Apêndice A.\n",
 "90_apendice.txt": "QUEBRA\n# Apêndice A: Prompts da pesquisa de mercado (Entrega 1)\n**Ferramenta e data:** Claude Code (Anthropic), em 05 e 06/10/2026, com subagentes autorizados a buscar e abrir páginas na web.\n**Método.** A IA foi usada na pesquisa e como apoio à redação, com revisão e reescrita final da equipe. Cinco frentes de pesquisa rodaram em paralelo (camadas, concorrência, dependências, defensabilidade e cálculos), com o mesmo contexto (os Quadros 3 a 6 transcritos) e as mesmas regras: não inventar números sobre a Lumis; não pesquisar as empresas fictícias e buscar, no lugar delas, análogos reais rotulados como tal; citar fonte aberta para todo fato externo; separar fato, dado do caso, cálculo e hipótese.\n**Verificação.** Um segundo agente tentou refutar cada achado, reabrindo a fonte e refazendo as contas. Dos 85 achados, 52 foram confirmados, 32 corrigidos e 1 descartado; os não confirmados ficaram fora do texto. Seguem os dois prompts de pesquisa mais usados e o de verificação.\n## Prompt 1: Concorrentes diretos, indiretos e potenciais\nCAIXA\n## Tarefa: concorrentes diretos, indiretos e potenciais (incluindo quem tem distribuição)\nLevante análogos REAIS no Brasil (e globais com presença no Brasil) para cada tipo de participante do Quadro 6, e acrescente tipos que o quadro não lista:\n- sistemas hospitalares/prontuário eletrônico (EHR/HIS) com IA embutida ou anunciada (ex.: verificar Epic, Oracle Health/Cerner, Philips Tasy, MV, TOTVS Saúde e outros com presença no Brasil): base instalada no Brasil, recursos de IA preditiva, triagem ou priorização anunciados, com fonte e data;\n- healthtechs brasileiras de IA clínica, triagem ou risco (confirme a existência e o produto de cada uma; não liste nomes sem fonte);\n- BI e analytics hospitalar; consultorias e integradores que fazem IA sob medida em saúde;\n- big techs e provedores de modelo com produtos de saúde (ex.: verificar ofertas de Microsoft, Google, Amazon, OpenAI, Anthropic voltadas à saúde, com data de lançamento);\n- concorrentes nos outros dois segmentos da Lumis (seguradoras: classificação de sinistros; bancos: risco de crédito), pois 14 das 38 contas estão fora da saúde (Quadro 3).\nPara cada um, classifique como DIRETO (mesmo problema, mesmo cliente), INDIRETO (resolve o problema de outro jeito) ou POTENCIAL (ainda não compete, mas tem distribuição nos mesmos clientes), e explique o mecanismo de ameaça. Destaque o padrão \"distribuição vence qualidade isolada\" (bundling no EHR) com evidência real, se existir.\nFIMCAIXA\n**Resultado obtido:** 17 achados a partir de 25 buscas. O principal é o padrão de distribuição embutida no prontuário, com análogos brasileiros (MV com 894 hospitais e Tasy com 500 na América Latina, segundo o KLAS; Tasy na Rede D'Or). Apareceram também a entrada de OpenAI e Anthropic na saúde (jan/2026), healthtechs e integradores nacionais e concorrentes em seguros e crédito.\n**Verificação:** 6 achados confirmados, 10 corrigidos e 1 descartado (um resultado de caso clínico atribuído à empresa errada). Uma crítica de completude apontou a falta dos grandes compradores que desenvolvem IA própria, e uma pesquisa complementar trouxe Einstein, Itaú, Bradesco e Porto. Antes de entrarem no texto, as fontes externas citadas foram reabertas em 05/10/2026, e os números citados foram confirmados.\n## Prompt 2: Dependências críticas de fornecedores\nCAIXA\n## Tarefa: dependências críticas e o que acontece se cada fornecedor mudar preço, termos ou escopo\nOs quatro fornecedores estruturais estão no Quadro 5 (modelo fundacional único, nuvem global, consórcio de bases clínicas, 38 contratos de dados de clientes). Para cada um, levante PRECEDENTES REAIS que mostram que o risco é concreto:\n- modelos fundacionais: casos documentados de mudança de preço, descontinuação/depreciação de modelos com prazo curto, mudanças de termos de uso ou de política de dados, e provedores lançando produto próprio para saúde (data e fonte);\n- nuvem: custos e prazos típicos de migração, taxas de saída de dados (egress) e mudanças regulatórias recentes sobre elas (ex.: EU Data Act), lock-in, regiões de dados no Brasil;\n- bases clínicas de referência licenciadas: como funcionam o licenciamento e os reajustes no setor (se não achar fonte, diga);\n- dados de clientes: o que a LGPD (dado de saúde é sensível, art. 11) e a ANPD dizem sobre reutilizar dados de clientes para treinar modelos; precedentes de disputa contratual ou regulatória sobre uso de dados de saúde para IA;\n- câmbio: volatilidade do real frente ao dólar nos últimos 3–5 anos (máximas e mínimas, com fonte, ex.: Banco Central), já que 72,2% do custo é em dólar.\nPara cada fornecedor, descreva qualitativamente o impacto no negócio da Lumis se ele (i) subir preço, (ii) mudar termos, (iii) mudar escopo ou virar concorrente. A quantificação fica a cargo de outra frente; aqui foque em evidência de que o evento é plausível.\nFIMCAIXA\n**Resultado obtido:** 20 achados a partir de 22 buscas. Os principais precedentes reais são a descontinuação de modelos com aviso curto, o aumento de 4 vezes no preço do Haiku em 2024, a concentração da nuvem e os custos de saída, os limites da LGPD para dado de saúde (art. 11) e a volatilidade do câmbio, com PTAX entre R$ 4,62 e R$ 6,21 em cinco anos (Banco Central).\n**Verificação:** 15 achados confirmados e 5 corrigidos. Uma pesquisa complementar sobre o poder de negociação com o fornecedor de modelo (alternativas e custo de troca) mudou a ordem de criticidade, e os dados de clientes passaram à frente do modelo. Os cálculos de sensibilidade (margem, câmbio, runway) foram refeitos a partir do Quadro 4, e a fonte do aumento de preço do Haiku foi reaberta.\n## Prompt 3: Verificação adversarial (usado em todas as frentes)\nCAIXA\nVocê é verificador independente e cético de uma due diligence. Sua função é REFUTAR, não confirmar. Na dúvida, marque \"nao_confirmado\".\n[bloco de contexto comum]\n## Achados a verificar: <lista de achados da frente>\n## Como verificar\n- fato_externo: abra a URL e confirme se a fonte diz exatamente o que a afirmação diz (número, data, autoria). Se a URL falhar, procure a mesma informação em outra fonte primária. Afirmação mais forte que a fonte = \"parcial\", com correção.\n- dado_anexo: confira contra os quadros, palavra por palavra.\n- calculo: refaça a conta. Qualquer divergência = \"refutado\", com o valor correto.\n- hipotese: avalie se está marcada como hipótese e se é razoável.\nVerifique TODOS os IDs.\nFIMCAIXA\n**O que a verificação mudou:** entre outras correções, evitou quatro erros factuais no texto: uma data de lançamento antecipada indevidamente (um produto previsto para 2027), uma comparação de \"quase 10 vezes\" que na verdade era de 8,3 vezes, um percentual de mercado aplicado ao total quando valia só para IA generativa e uma data de fim de caixa calculada errado.\n"
}

CONTEUDO_MEMO = "TITULO: Memorando ao conselho\nSUBTITULO: Proposta do Vetor Capital: direção, riscos e condições prévias\nLINHA: **Para:** Conselho da Lumis Intelligence\nLINHA: **De:** Head of AI Management, pela equipe: Bruno Müller, Diego Franca Evangelista, Felipe Alef, Gustavo Halfen Simon e Maria Fernanda Barros\nLINHA: **Data:** 08/10/2026\n\nROTULO: Recomendação\n> Recomendamos aceitar o aporte do Vetor Capital com liberação por etapas, ligada a condições verificáveis: primeiro tornar o ativo defensável, crescer onde a Lumis já tem cliente e contratar em ondas; novos países e plataforma, só depois. Se o Vetor exigir países ou plataforma em produção antes das condições, recomendamos não aceitar o aporte nesses termos.\n\nA proposta do Vetor Capital, R$ 120 milhões por 22%, avalia a Lumis em R$ 425,5 milhões antes do aporte (cerca de 10,3 vezes a receita recorrente anual) e vale por 60 dias [fonte: Cap. 2, texto após o Quadro 3]. Não aceitar tem custo: no fechamento do 2º trimestre de 2026, a Lumis tinha caixa para 11,6 meses [fonte: Cap. 2, Quadro 3], e outra fonte de capital [não consta].\n\n## Direção proposta\n\nHoje o ativo da Lumis não passa numa auditoria independente. Dos dados de clientes, 35,4% têm autorização fraca ou nenhuma para treino, e o modelo erra 31,8% dos idosos de CEP D/E que precisavam de prioridade, três vezes o melhor grupo [fonte: Cap. 2, Quadros 7 e 10].\n\nO ativo a construir é o dado de desfecho com direito de uso limpo e a prova auditável de desempenho por grupo. O crescimento vem de hospitais e seguradoras médios no Brasil, que, sem IA própria, dependem mais de fornecedores externos [hipótese]; a Aster não chega a seguradoras e bancos, 14 das 38 contas [fonte: Cap. 2, Quadros 3 e 6]. O novo módulo de risco de crédito para os cinco bancos já clientes é exceção ao foco, só até o modo sombra (em paralelo, sem decidir).\n\nTAB[memo_aporte]: 4.6;11.4 | Uso do aporte. Fonte: Cap. 2, 1.1 (proposta do Vetor); recomendação da equipe.\n| O Vetor propõe | Recomendamos |\n| Triplicar o time técnico | Contratar em ondas, depois que a regra única de prova e a mesa de campo (condições 2 e 4) estiverem rodando; revisão em 01/2027 |\n| Abrir dois novos países | Nenhum nesta janela de seis meses; o México é reavaliado ao fim dela, e o segundo país [não consta] |\n| Transformar em plataforma | Só 1,0 mês-pessoa, para testar se a prova de desempenho por grupo vira oferta |\n\n## Riscos que assumimos\n\nTAB[memo_riscos]: 9.2;2.8;4.0 | Riscos assumidos. Fonte: Cap. 2, Anexo A; LGPD, art. 52, II; contas da equipe.\n| Risco | Quem responde | Resposta |\n| A Aster ocupa os hospitais enquanto 78% da capacidade do semestre vai para melhoria | CEO | Crédito por etapas e teste da prova por grupo; revisão ao fim da janela |\n| As três recusas (veterinário, agente de triagem e México) deixam de lado 57,4% da receita potencial listada (R$ 11,3 mi), sem receita nova garantida na janela | CEO | Revisão ao fim da janela |\n| A restrição devolve cerca de 154 mil priorizações por mês à triagem dos hospitais [hipótese: distribuição igual à da base]; carga por hospital e receita afetada [não consta] | Head of AI Management e CTO | O modelo em paralelo mede a correção; religar exige os dois cargos |\n| Se a renegociação falhar, a Lumis retira 35,4% dos dados de clientes e retreina; multa da LGPD de até 2% do faturamento, limitada a R$ 50 mi por infração | DPO | Rastrear a origem de cada registro para isolar os frágeis |\n| Triplicar o time antes das intervenções leva gente nova para o mesmo pressuposto [hipótese]; o Responsável por Dados está em quase todos os portões, num time com rotatividade de 27% | CEO e CTO | Contratar em ondas, com revisão em 01/2027 |\n\n## Condições prévias\n\nEm até 30 dias, antes de fechar com o Vetor:\n1. Restringir já a recomendação automática para idosos de CEP C e D/E, com o modelo em paralelo, como prevê o compromisso de Não Amplificação de Danos da Fase 1 (Head of AI Management ordena; CTO executa).\n2. Retirar ou tornar auditáveis as cinco afirmações do painel comercial pela regra única de prova: nenhum número de desempenho sai da empresa sem o resultado de campo por grupo (Head of AI Management; o Responsável Comercial aplica).\n3. Aprovar a linha de responsabilidade (CEO) e testar a suspensão por cliente e por grupo, com teto de 48 horas e novo teste a cada trimestre (Head of AI Management conduz; CTO executa).\n4. Instalar a mesa de campo, reunião quinzenal que lê o erro por grupo cruzado com as reclamações, com primeira reunião em 15 dias (Responsável por Produto; o cruzamento é do Responsável por Dados).\n5. Obter o parecer da DPO sobre os seis instrumentos de dados antes do próximo treino; registro sem base fica fora do treino.\n\nAntes do vencimento dos contratos:\n6. Assinar aditivos de direito de uso com a Seguradora Prisma (antes de 12/2026) e o Hospital Vila Ipê (antes de 03/2027), incluindo o histórico (DPO).\n\nAntes de cada frente nova (país, plataforma ou novo módulo de crédito em produção), a CEO decide se ela entra, com parecer obrigatório do Head of AI Management e da DPO, e exige:\n7. política de privacidade para dados de saúde aprovada (DPO; previsão de aprovação [não consta]);\n8. erro por grupo no Brasil até 2 vezes o do melhor e já publicado na métrica pública;\n9. validação local por grupo, com amostra suficiente, e lei de dados do país conferida pela DPO;\n10. relatório de impacto à proteção de dados (LGPD, art. 38) da plataforma e de cada país;\n11. no novo módulo de crédito, dado de saúde fora do modelo e aprovação por faixa de CEP de pelo menos 0,8 da melhor;\n12. modo sombra até quatro leituras quinzenais com o pior grupo até 1,5 vez o melhor, responsável do domínio nomeado e suspensão testada no novo mercado.\n\n## O que pedimos ao conselho\n\n1. Aprovar a direção e a cláusula de saída para negociar com o Vetor.\n2. Endossar a restrição imediata e as três recusas.\n\nNOTA: Detalhe e fontes no dossiê da Fase 2. Metas, limites, prazos e cargos são propostas da equipe sobre o Cap. 2, Anexo A.\n"


if __name__ == '__main__':
    modo = sys.argv[1] if len(sys.argv) > 1 else 'lint'
    if modo == 'lint':
        avisos, figs, tabs, refs = montar('dossie', ':dossie', None)
        avisos += montar('memo', ':memo', None)[0]
    else:
        avisos, figs, tabs, refs = montar(modo, ':dossie' if modo == 'dossie' else ':memo', sys.argv[2])
    for a in avisos:
        print('AVISO', a)
    print(f'avisos: {len(avisos)}')
